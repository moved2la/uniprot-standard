"""Enumerate the candidate pool for each tier by querying UniProt, by code.

Inputs
  data/gene-ontology/resolved_terms.ini     (from resolve_ontology_term.py)
  data/gene-ontology/tier<N>_subtree.tsv
  data/gene-ontology/tier<N>_neighbours.tsv

Definition of the pool (per tier)
  Human (organism_id:9606), reviewed (reviewed:true) UniProtKB entries annotated to
  the tier's term or to any term in its descendant subtree. The pool is computed as
  the UNION of one query per subtree term, so its definition does not depend on how
  UniProt expands a term. UniProt's own term-level query is run as a check and the
  two are compared; a difference is recorded, never silently absorbed.

Sensitivity
  The same term-level query is run for each immediate neighbour of the tier's term
  (next broader, next narrower) and the set differences against the pool are
  recorded. This is the sensitivity of the pool to the category boundary.

Measured tier (D56)
  A tier whose decisions section says `definition = measured_remainder` is not a Gene
  Ontology category. Its pool is every gene quantified in the primary Layer B dataset
  (the [source.*] named in the decisions section, read through its [columns.*] map in
  config/mass_fraction_decisions.ini and data/literature/manifest.ini) that maps by
  exact gene symbol to exactly one reviewed human UniProt entry, is not already in an
  ontology tier's pool, and is not in the cited contaminant list. Every gene's outcome
  is written to data/tier<N>_gene_mapping.tsv. Rules M1-M5 (docs/conventions.md).
  The ontology tiers are enumerated first, so "already in an ontology tier" is known.

Pools by published accession (Step 7b, D107; rules P1-P6)
  `--category <name>` enumerates the pools a category declares as [pool.<name>] sections in
  config/<category>/protein_set_decisions.ini. Each pool is the quantified set of its primary
  dataset (D98): no Gene Ontology seed, no tier split, no gene-symbol lookup. The dataset's
  identity column holds MaxQuant protein groups -- semicolon-joined accessions in the order the
  authors published -- and the pool entry of a row is the first accession of its group that is a
  reviewed human entry. Keratins (D98) are recognised by UniProt's keyword `Keratin` on the entry;
  the compartment's named contaminants by the dataset's own gene cell. Nothing is excluded
  silently: every row's outcome is written, and the excluded rows are listed with their share.

Outputs
  data/tier<N>_pool.tsv
  data/tier<N>_pool_sensitivity.tsv
  data/pool_queries.ini
  outputs/pool_overlap.tsv         (accessions appearing in more than one tier)
  -- with --category <name> --
  data/<category>/pool_<pool>.tsv                    the pool: one row per entry
  data/<category>/<pool>_dataset_rows.tsv            every dataset row with its outcome (P1-P5)
  data/<category>/<pool>_rows_excluded_with_share.tsv  what the rules removed, largest first
  data/<category>/pool_queries.ini                   the run record
  outputs/<category>/intermediate/protein_set/pool_overlap.tsv
"""

from __future__ import annotations

import argparse
import sys

import requests

from pipeline import common
from pipeline.flags import append_flag, reset_flags
from pipeline.uniprot_client import UniProtClient

FIELDS = ["accession", "gene_primary", "protein_name", "length", "reviewed", "organism_id"]
STAGE = "enumerate_pool"


def term_query(term_id: str) -> str:
    """UniProt query-field syntax uses the numeric part of a Gene Ontology ID."""
    num = term_id.split(":", 1)[1]
    return f"(organism_id:9606) AND (reviewed:true) AND (go:{num})"


def _col(row: dict[str, str], *candidates: str) -> str:
    for c in candidates:
        if c in row:
            return row[c]
    return ""


def _accession(row: dict[str, str]) -> str:
    return _col(row, "Entry", "accession") or next(iter(row.values()))


def enumerate_tier(client: UniProtClient, tier: str, term_id: str, term_name: str,
                   subtree: list[dict[str, str]], neighbours: list[dict[str, str]], log):
    # --- definitional pool: union of one query per subtree term ---------------
    members: dict[str, dict[str, str]] = {}
    returned_by: dict[str, list[str]] = {}
    per_term_counts: list[list] = []
    for t in subtree:
        tid = t["term_id"]
        rows = client.search_tsv(term_query(tid), FIELDS)
        per_term_counts.append(["subtree", tid, t["name"], t["relation_to_parent"], t["depth"], len(rows)])
        for r in rows:
            acc = _accession(r)
            members.setdefault(acc, r)
            returned_by.setdefault(acc, []).append(tid)
    pool = set(members)
    log.info("tier %s: union over %d subtree terms -> %d entries", tier, len(subtree), len(pool))

    # --- check: UniProt's own expansion of the tier term ------------------------
    top_rows = client.search_tsv(term_query(term_id), FIELDS)
    top = {_accession(r) for r in top_rows}
    only_in_top = sorted(top - pool)
    only_in_union = sorted(pool - top)
    expansion_matches = not only_in_top and not only_in_union
    log.info("tier %s: term-level query -> %d entries; matches union: %s "
             "(only_in_term_query=%d, only_in_union=%d)",
             tier, len(top), expansion_matches, len(only_in_top), len(only_in_union))
    if not expansion_matches:
        append_flag(common.FLAGS_TSV, stage=STAGE, subject=f"tier.{tier}.expansion",
                    rule="descendant_expansion_check",
                    detail=(f"term-level query and per-term union differ: "
                            f"only_in_term_query={only_in_top} only_in_union={only_in_union}"))

    # --- sensitivity: immediate neighbours ---------------------------------------
    sens_rows: list[list] = []
    sens_rows.append(["tier_term", term_id, term_name, "", "term-level query", len(top),
                      len(only_in_top), len(only_in_union), ";".join(only_in_top), ";".join(only_in_union)])
    for nb in neighbours:
        tid = nb["term_id"]
        rows = client.search_tsv(term_query(tid), FIELDS)
        s = {_accession(r) for r in rows}
        extra = sorted(s - pool)
        missing = sorted(pool - s)
        sens_rows.append([nb["direction"], tid, nb["name"], nb["relation"], "term-level query",
                          len(rows), len(extra), len(missing), ";".join(extra), ";".join(missing)])
        log.info("tier %s: %s neighbour %s (%s): %d entries, %d not in pool, %d of pool absent",
                 tier, nb["direction"], tid, nb["name"], len(rows), len(extra), len(missing))

    # --- write --------------------------------------------------------------------
    pool_rows = []
    for acc in sorted(pool):
        r = members[acc]
        pool_rows.append([
            acc,
            _col(r, "Gene Names (primary)", "gene_primary"),
            _col(r, "Protein names", "protein_name"),
            _col(r, "Length", "length"),
            _col(r, "Reviewed", "reviewed"),
            _col(r, "Organism (ID)", "organism_id"),
            ";".join(returned_by[acc]),
            "true" if acc in top else "false",
        ])
    common.write_tsv(common.DATA_DIR / f"tier{tier}_pool.tsv",
                     ["accession", "gene_primary", "protein_name", "length", "reviewed", "organism_id",
                      "subtree_terms_returning_entry", "in_term_level_query"],
                     pool_rows)
    common.write_tsv(common.DATA_DIR / f"tier{tier}_pool_sensitivity.tsv",
                     ["query_label", "term_id", "term_name", "relation", "query_kind", "n_entries",
                      "n_not_in_pool", "n_pool_not_returned", "accessions_not_in_pool", "pool_accessions_not_returned"],
                     sens_rows)
    common.write_tsv(common.DATA_DIR / f"tier{tier}_pool_per_term_counts.tsv",
                     ["kind", "term_id", "term_name", "relation_to_parent", "depth", "n_entries"],
                     per_term_counts)
    return pool, len(top), expansion_matches


# ----------------------------------------------------------------------------
# measured remainder (D56)
# ----------------------------------------------------------------------------

MEASURED_FIELDS = ["accession", "gene_primary", "gene_synonym", "protein_name", "length", "reviewed", "organism_id"]
GENE_BATCH = 50


def dataset_gene_cells(source_id: str, file_key: str, log) -> list[str]:
    """Every value of the identity column of the primary dataset, as read (rules B0, I1)."""
    import io
    from openpyxl import load_workbook
    dec = common.read_ini(common.CONFIG_DIR / "mass_fraction_decisions.ini")
    manifest = common.read_ini(common.literature_manifest())
    cols = dec[f"columns.{source_id}.{file_key}"]
    m = manifest[f"file.{source_id}.{file_key}"]
    path = common.REPO_ROOT / m["path"]
    got = common.sha256_file(path)
    if got != m["sha256"].lower():
        raise SystemExit(f"[STOP] {m['path']} sha256 {got} != manifest {m['sha256']}")
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    ws = wb[cols["sheet"]]
    header_row = int(cols["header_row"])
    cells = []
    for r_index, row in enumerate(ws.iter_rows(values_only=True), 1):
        if r_index < header_row:
            continue
        if r_index == header_row:
            headers = ["" if v is None else str(v).strip() for v in row]
            hits = [j for j, h in enumerate(headers) if h == cols["identity_column"].strip()]
            if len(hits) != 1:
                raise SystemExit(f"[STOP] identity column {cols['identity_column']!r} matched {len(hits)} headers")
            j = hits[0]
            continue
        v = row[j] if j < len(row) else None
        if v is not None and str(v).strip():
            cells.append(str(v).strip())
    wb.close()
    log.info("measured tier: %d identity cells read from %s sha256 %s", len(cells), m["path"], got)
    return cells


def contaminant_accessions(source_id: str, log) -> set[str]:
    """Accessions in the cited contaminant FASTA (gzip or plain). The header's first
    token, with a leading 'CON__' or 'sp|'/'tr|' prefix removed, is the accession."""
    import gzip
    manifest = common.read_ini(common.literature_manifest())
    m = manifest[f"file.{source_id}.file.1"]
    path = common.REPO_ROOT / m["path"]
    got = common.sha256_file(path)
    if got != m["sha256"].lower():
        raise SystemExit(f"[STOP] {m['path']} sha256 {got} != manifest {m['sha256']}")
    data = path.read_bytes()
    text = (gzip.decompress(data) if data[:2] == b"\x1f\x8b" else data).decode("utf-8", errors="replace")
    accs = set()
    for line in text.splitlines():
        if not line.startswith(">"):
            continue
        tok = line[1:].split()[0] if line[1:].split() else ""
        tok = tok.replace("CON__", "")
        if "|" in tok:
            parts = tok.split("|")
            tok = parts[1] if len(parts) > 1 else parts[0]
        accs.add(tok.split("-")[0])
    log.info("measured tier: %d contaminant accessions parsed from %s sha256 %s", len(accs), m["path"], got)
    return accs


def enumerate_measured_tier(client: UniProtClient, tier: str, sec, ontology_pools: dict[str, set[str]], log):
    """Rules
    M1  genes = every identity cell of the primary dataset, split on ';', stripped, unique.
    M2  each gene is looked up with (organism_id:9606) AND (reviewed:true) AND (gene_exact:<gene>),
        in batches; a returned entry belongs to a gene when the gene equals its primary
        symbol or one of its synonyms (case-sensitive, as UniProt reports them).
    M3  exactly one entry -> candidate. Zero -> unmapped, listed. Several -> ambiguous, listed, excluded.
        A symbol UniProt's parser rejects even when quoted -> query_rejected, listed.
    M4  a candidate already in an ontology tier's pool -> in_tier<N>, not a member.
    M5  a candidate in the contaminant list -> contaminant, not a member. Everything else -> member.
    Nothing is named; nothing is filtered by abundance."""
    cells = dataset_gene_cells(sec["dataset_source"], sec["dataset_file"], log)
    genes = sorted({g.strip() for c in cells for g in c.split(";") if g.strip()})
    contaminants = contaminant_accessions(sec["contaminant_source"], log)
    in_ontology = {acc: t for t, accs in ontology_pools.items() for acc in accs}
    found: dict[str, list[dict[str, str]]] = {g: [] for g in genes}
    rejected: set[str] = set()

    def quoted(g: str) -> str:
        return 'gene_exact:"' + g.replace('"', '') + '"'      # quoted: symbols may carry ':' '.' '-' (M2)

    def attribute(rows, batch):
        for r in rows:
            primary = _col(r, "Gene Names (primary)", "gene_primary").strip()
            syn = _col(r, "Gene Names (synonym)", "gene_synonym")
            names = {primary} | {s.strip() for s in syn.split() if s.strip()}
            for g in batch:
                if g in names:
                    found[g].append(dict(r, _matched_by="primary" if g == primary else "synonym"))

    for i in range(0, len(genes), GENE_BATCH):
        batch = genes[i:i + GENE_BATCH]
        q = "(organism_id:9606) AND (reviewed:true) AND (" + " OR ".join(quoted(g) for g in batch) + ")"
        try:
            rows = client.search_tsv(q, MEASURED_FIELDS)
            attribute(rows, batch)
            n = len(rows)
        except requests.exceptions.HTTPError as exc:                # a symbol UniProt cannot parse: retry one by one
            log.warning("measured tier: batch %d-%d rejected (%s); querying each gene alone", i + 1, i + len(batch), exc)
            n = 0
            for g in batch:
                try:
                    rows = client.search_tsv("(organism_id:9606) AND (reviewed:true) AND " + quoted(g), MEASURED_FIELDS)
                    attribute(rows, [g])
                    n += len(rows)
                except requests.exceptions.HTTPError as exc2:
                    rejected.add(g)
                    log.warning("measured tier: query rejected for %r (%s)", g, exc2)
        log.info("measured tier: genes %d-%d of %d queried; %d rows", i + 1, min(i + GENE_BATCH, len(genes)), len(genes), n)
    members: dict[str, dict[str, str]] = {}
    mapping_rows = []
    counts = {"member": 0, "unmapped": 0, "ambiguous": 0, "contaminant": 0, "query_rejected": 0}
    for g in genes:
        rows = found[g]
        accs = sorted({_accession(r) for r in rows})
        if g in rejected:
            outcome = "query_rejected"
        elif not accs:
            outcome = "unmapped"
        elif len(accs) > 1:
            outcome = "ambiguous"
        else:
            acc = accs[0]
            if acc in in_ontology:
                outcome = f"in_tier{in_ontology[acc]}"
            elif acc in contaminants:
                outcome = "contaminant"
            else:
                outcome = "member"
                members[acc] = rows[0]
        counts[outcome] = counts.get(outcome, 0) + 1
        mapping_rows.append([g, outcome, ";".join(accs), rows[0]["_matched_by"] if len(accs) == 1 else "",
                             _col(rows[0], "Gene Names (primary)", "gene_primary") if len(accs) == 1 else ""])
    common.write_tsv(common.DATA_DIR / f"tier{tier}_gene_mapping.tsv",
                     ["gene", "outcome", "accessions", "matched_by", "uniprot_primary_symbol"], mapping_rows)
    pool_rows = []
    for acc in sorted(members):
        r = members[acc]
        pool_rows.append([acc, _col(r, "Gene Names (primary)", "gene_primary"), _col(r, "Protein names", "protein_name"),
                          _col(r, "Length", "length"), _col(r, "Reviewed", "reviewed"), _col(r, "Organism (ID)", "organism_id"),
                          "", ""])
    common.write_tsv(common.DATA_DIR / f"tier{tier}_pool.tsv",
                     ["accession", "gene_primary", "protein_name", "length", "reviewed", "organism_id",
                      "subtree_terms_returning_entry", "in_term_level_query"], pool_rows)
    log.info("measured tier %s: %d genes -> %s", tier, len(genes), counts)
    return set(members), counts


# ----------------------------------------------------------------------------
# pools by published accession (D107; rules P1-P6, docs/conventions.md)
# ----------------------------------------------------------------------------

LOOKUP_FIELDS = ["accession", "reviewed", "organism_id", "gene_primary", "protein_name", "length", "keywordid"]
HUMAN_TAXON = "9606"
CONTAMINANT_PREFIX = "CON__"
DECOY_PREFIX = "REV__"


def read_sheet_by_letter(path, sheet: str, header_row: int, letters: dict[str, str], log) -> list[dict]:
    """Every data row below `header_row` of `sheet`, the wanted cells keyed by name, read by
    COLUMN LETTER (the blood tables carry duplicate and stray header names, so a header string
    cannot identify a column). Returns [{"row_index": int, <key>: cell or None, ...}]."""
    import io
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"[STOP] sheet {sheet!r} not in {path.name}: {wb.sheetnames}")
    ws = wb[sheet]
    idx = {k: column_index_from_string(v.strip()) - 1 for k, v in letters.items()}
    rows = []
    for r_index, row in enumerate(ws.iter_rows(values_only=True), 1):
        if r_index <= header_row:
            continue
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec = {"row_index": r_index}
        for k, j in idx.items():
            rec[k] = row[j] if j < len(row) else None
        rows.append(rec)
    wb.close()
    log.info("%s sheet %r: %d non-empty rows below row %d; columns %s",
             path.name, sheet, len(rows), header_row, ", ".join(f"{k}={v}" for k, v in letters.items()))
    return rows


def split_letters(value: str) -> list[str]:
    return [x.strip() for x in value.replace(";", ",").split(",") if x.strip()]


def classify_token(token: str) -> tuple[str, str]:
    """P2: a group token -> (kind, canonical accession or ''). kinds: contaminant_marker, decoy,
    accession (isoform suffix stripped, R1), blank."""
    t = token.strip()
    if not t:
        return "blank", ""
    if t.startswith(CONTAMINANT_PREFIX):
        return "contaminant_marker", ""
    if t.startswith(DECOY_PREFIX):
        return "decoy", ""
    return "accession", t.split("-", 1)[0]


def lookup_accessions(client: UniProtClient, tokens: list[str], log) -> dict[str, dict[str, str]]:
    """P3: what UniProt says about each accession: reviewed, organism, primary symbol, keyword ids.
    Batches of GENE_BATCH `accession:` terms; a row is attributed to a token when the token equals
    the entry's primary accession. A token with no such row is not returned (obsolete, or now a
    secondary accession of another entry): it is recorded as such, never guessed."""
    found: dict[str, dict[str, str]] = {}
    todo = sorted(set(tokens))
    for i in range(0, len(todo), GENE_BATCH):
        batch = todo[i:i + GENE_BATCH]
        q = "(" + " OR ".join(f"accession:{t}" for t in batch) + ")"
        rows = client.search_tsv(q, LOOKUP_FIELDS)
        for r in rows:
            acc = _accession(r)
            if acc in batch:
                found[acc] = {
                    "reviewed": _col(r, "Reviewed", "reviewed"),
                    "organism_id": _col(r, "Organism (ID)", "organism_id"),
                    "gene_primary": _col(r, "Gene Names (primary)", "gene_primary"),
                    "protein_name": _col(r, "Protein names", "protein_name"),
                    "length": _col(r, "Length", "length"),
                    "keyword_ids": _col(r, "Keyword ID", "keywordid"),
                }
        log.info("accession lookup: %d-%d of %d; %d rows", i + 1, min(i + GENE_BATCH, len(todo)), len(todo), len(rows))
    return found


GENE_FALLBACK_FIELDS = MEASURED_FIELDS + ["keywordid"]


def _info_of(r: dict[str, str]) -> dict[str, str]:
    return {
        "reviewed": _col(r, "Reviewed", "reviewed"),
        "organism_id": _col(r, "Organism (ID)", "organism_id"),
        "gene_primary": _col(r, "Gene Names (primary)", "gene_primary"),
        "protein_name": _col(r, "Protein names", "protein_name"),
        "length": _col(r, "Length", "length"),
        "keyword_ids": _col(r, "Keyword ID", "keywordid"),
    }


def lookup_secondary(client: UniProtClient, tokens: list[str], log) -> dict[str, dict[str, dict[str, str]]]:
    """P3b (D111): tokens UniProt no longer returns as an active primary accession are asked for ONE AT A
    TIME with the `sec_acc` query field (secondary accession). UniProt has no return field for secondary
    accessions, so attribution is by the query itself, and every hit is then confirmed against the
    entry's own `secondaryAccessions` list in its JSON. Returns token -> {entry accession: info}; a
    token with no hit is obsolete (deleted) -- nothing carries it now.
    If UniProt refuses the query field, P3b resolves nothing this run (logged), and P4b carries on."""
    out: dict[str, dict[str, dict[str, str]]] = {}
    todo = sorted(set(tokens))
    for n, t in enumerate(todo, 1):
        try:
            rows = client.search_tsv(f"sec_acc:{t}", LOOKUP_FIELDS)
        except requests.exceptions.HTTPError as exc:
            log.warning("secondary-accession lookup: UniProt refused the sec_acc query field (%s); "
                        "P3b resolves nothing this run, the gene fallback (P4b) carries on", exc)
            return out
        for r in rows:
            acc = _accession(r)
            info = _info_of(r)
            if not is_reviewed_human(info):
                continue
            try:
                secs = set(client.entry_json(acc).get("secondaryAccessions", []))
            except requests.RequestException as exc:                      # confirmation failed: not attributed
                log.warning("secondary-accession lookup: %s -> %s could not be confirmed (%s)", t, acc, exc)
                continue
            if t in secs:
                out.setdefault(t, {})[acc] = info
            else:
                log.info("secondary-accession lookup: %s returned for sec_acc:%s but does not list it; ignored", acc, t)
        if n % 25 == 0 or n == len(todo):
            log.info("secondary-accession lookup: %d of %d tokens asked; %d resolved so far", n, len(todo), len(out))
    return out


def lookup_genes(client: UniProtClient, genes: list[str], log) -> dict[str, dict[str, dict[str, str]]]:
    """P4b (D111): the muscle M2 lookup, unchanged in its query and its attribution -- a returned entry
    belongs to a gene when the gene equals its primary symbol or one of its synonyms -- used here only
    for a group with no reviewed human accession. Returns gene -> {entry accession: info}."""
    out: dict[str, dict[str, dict[str, str]]] = {g: {} for g in genes}
    todo = sorted(set(genes))

    def quoted(g: str) -> str:
        return 'gene_exact:"' + g.replace('"', '') + '"'

    def attribute(rows, batch):
        for r in rows:
            primaries = {x.strip() for x in _col(r, "Gene Names (primary)", "gene_primary").split(";") if x.strip()}
            syn = _col(r, "Gene Names (synonym)", "gene_synonym")
            synonyms = {x.strip() for x in syn.split() if x.strip()}
            for g in batch:
                if g in primaries or g in synonyms:
                    out[g][_accession(r)] = dict(_info_of(r), _matched_by="primary" if g in primaries else "synonym")

    for i in range(0, len(todo), GENE_BATCH):
        batch = todo[i:i + GENE_BATCH]
        q = "(organism_id:9606) AND (reviewed:true) AND (" + " OR ".join(quoted(g) for g in batch) + ")"
        try:
            rows = client.search_tsv(q, GENE_FALLBACK_FIELDS)
            attribute(rows, batch)
        except requests.exceptions.HTTPError as exc:
            log.warning("gene fallback: batch %d-%d rejected (%s); querying each gene alone", i + 1, i + len(batch), exc)
            for g in batch:
                try:
                    attribute(client.search_tsv("(organism_id:9606) AND (reviewed:true) AND " + quoted(g), GENE_FALLBACK_FIELDS), [g])
                except requests.exceptions.HTTPError as exc2:
                    log.warning("gene fallback: query rejected for %r (%s)", g, exc2)
        log.info("gene fallback lookup: %d-%d of %d", i + 1, min(i + GENE_BATCH, len(todo)), len(todo))
    return out


def is_reviewed_human(info: dict[str, str] | None) -> bool:
    return bool(info) and info["reviewed"].strip().lower() == "reviewed" and info["organism_id"].strip() == HUMAN_TAXON


def is_inactive(info: dict[str, str] | None) -> bool:
    """UniProt answers a search for a deleted or demerged accession with an INACTIVE stub: the accession,
    no reviewed status, no organism. Such a token has no entry; it goes to the secondary-accession lookup (P3b)."""
    return bool(info) and not info["reviewed"].strip() and not info["organism_id"].strip()


def has_keyword(info: dict[str, str], keyword_id: str) -> bool:
    return keyword_id in {k.strip() for k in info.get("keyword_ids", "").split(";") if k.strip()}


def gene_tokens(cell) -> list[str]:
    return [g.strip() for g in str(cell or "").split(";") if g.strip()]


def published_share(rec: dict, cols) -> float | None:
    """The listing's share of a row, from the columns the map names, before any join: the mean over
    `share_columns` (a per-replicate mass share as published) or 10**`abundance_column` when the
    map says `abundance_scale = log10`. Preliminary -- the weights come from the mass-fraction
    stage; this number exists so the excluded-rows table can say what the rules removed."""
    if cols.get("share_columns"):
        vals = []
        for k in range(len(split_letters(cols["share_columns"]))):
            v = rec.get(f"share_{k}")
            if v is None or str(v).strip() == "":
                continue
            try:
                vals.append(float(v))
            except ValueError:
                continue
        return (sum(vals) / len(vals)) if vals else None
    if cols.get("abundance_column"):
        v = rec.get("abundance")
        if v is None or str(v).strip() == "":
            return None
        try:
            x = float(v)
        except ValueError:
            return None
        return 10 ** x if cols.get("abundance_scale", "").strip().lower() == "log10" else x
    return None


def enumerate_published_pool(client: UniProtClient, category: str, pool: str, sec, rules, log) -> tuple[set[str], dict]:
    """Rules
    P1  rows = every non-empty data row of the primary dataset's identity column, in file order.
    P2  a row's group = its identity cell split on ';' in the published order; each token is a
        CON__ contaminant marker, a REV__ decoy, or an accession with any isoform suffix removed (R1).
    P3  every accession token is looked up by `accession:`; reviewed, organism and keyword ids are read.
    P3b (D111) a token not returned as a primary accession, or returned as an INACTIVE stub (deleted or
        demerged: no reviewed status, no organism), is asked for again with the secondary-accession field:
        exactly one reviewed human entry lists it -> that entry (member_via_secondary_accession);
        several (a demerged entry) -> all of them, the row is SHARED between them and the mass-fraction
        join splits it (D48; member_via_secondary_accession_shared); none -> the token is obsolete.
    P4  the row's entry = the first token in group order that is a reviewed human entry. A group whose
        FIRST token is a contaminant marker -> contaminant_group.
    P4b (D111) a group with no usable token: EACH gene of its gene cell is looked up by the muscle M2 query
        and attribution; M3 per gene -- exactly one reviewed human entry -> that entry; M3b (D111a): when
        several answer, one of which by its primary symbol and the rest by synonym, the primary-symbol
        entry stands; otherwise the gene is ambiguous. One gene -> member_via_gene; several genes, each
        resolved -> member_via_gene_shared (a shared row, split by D48); any gene unresolved -> the row is
        no_reviewed_human_entry, excluded and listed with its share.
    P5  an entry carrying the keratin keyword -> keratin; a row whose gene cell has a token in the pool's
        named contaminant genes -> named_contaminant. Both excluded, listed with their share.
    P6  members are the distinct entries of the surviving rows; a later row choosing an entry already
        chosen is member_duplicate_entry (its values are summed by the mass-fraction join, D47).
    Nothing is named in code: the sheet, the columns, the keyword id, and the gene names come from config."""
    dec = common.read_ini(common.category_config_dir(category) / "mass_fraction_decisions.ini")
    manifest = common.read_ini(common.literature_manifest())
    src, fkey = sec["dataset_source"], sec["dataset_file"]
    cols = dec[f"columns.{src}.{fkey}"]
    m = manifest[f"file.{src}.{fkey}"]
    path = common.REPO_ROOT / m["path"]
    got = common.sha256_file(path)
    if got != m["sha256"].lower():
        raise SystemExit(f"[STOP] {m['path']} sha256 {got} != manifest {m['sha256']}")
    if cols.get("identity_kind", "").strip() != "accession_list":
        raise SystemExit(f"[STOP] [columns.{src}.{fkey}] identity_kind must be accession_list for a published-accession pool")

    letters = {"identity": cols["identity_column"], "gene": cols["gene_column"]}
    for k, letter in enumerate(split_letters(cols.get("share_columns", ""))):
        letters[f"share_{k}"] = letter
    if cols.get("abundance_column"):
        letters["abundance"] = cols["abundance_column"]
    rows = read_sheet_by_letter(path, cols["sheet"], int(cols["header_row"]), letters, log)
    log.info("pool %s: %s %s (%s) sha256 %s; %d rows", pool, src, fkey, m["path"], got, len(rows))

    keratin_kw = rules["keratin_keyword_id"].strip()
    named = {g.strip() for g in sec.get("named_contaminant_genes", "").split(",") if g.strip()}

    # P1/P2: groups
    groups: list[tuple[dict, list[tuple[str, str]]]] = []
    for rec in rows:
        cell = "" if rec["identity"] is None else str(rec["identity"]).strip()
        toks = [classify_token(t) for t in cell.split(";")] if cell else []
        groups.append((rec, toks))

    # P3: look up first-accession tokens, then the rest of the groups that still need a token
    first_tokens = [next((acc for kind, acc in toks if kind == "accession"), "") for _, toks in groups]
    asked = {t for t in first_tokens if t}
    info = lookup_accessions(client, sorted(asked), log)
    remaining = set()
    for (rec, toks), first in zip(groups, first_tokens):
        if first and not is_reviewed_human(info.get(first)):
            remaining.update(acc for kind, acc in toks if kind == "accession" and acc not in asked)
    if remaining:
        info.update(lookup_accessions(client, sorted(remaining), log))

    # P3b: tokens no entry answers to by primary accession -> the secondary-accession lookup
    def first_usable(toks) -> str:
        return next((acc for kind, acc in toks if kind == "accession" and is_reviewed_human(info.get(acc))), "")
    unresolved = sorted({acc for _, toks in groups if not first_usable(toks)
                         for kind, acc in toks if kind == "accession" and (acc not in info or is_inactive(info[acc]))})
    secondary = lookup_secondary(client, unresolved, log) if unresolved else {}
    secondary = {t: {a: i for a, i in ents.items() if is_reviewed_human(i)} for t, ents in secondary.items()}
    secondary = {t: ents for t, ents in secondary.items() if ents}

    # P4b: groups still without an entry -> the gene cell, through the M2/M3 rules
    def secondary_hit(toks):
        return next((secondary[acc] for kind, acc in toks if kind == "accession" and acc in secondary), None)
    fallback_genes = sorted({g for rec, toks in groups
                             if toks and toks[0][0] != "contaminant_marker" and not first_usable(toks) and secondary_hit(toks) is None
                             for g in gene_tokens(rec["gene"])})
    by_gene = lookup_genes(client, fallback_genes, log) if fallback_genes else {}

    # P4-P6: outcomes
    counts: dict[str, int] = {}
    members: dict[str, dict] = {}
    row_table: list[list] = []
    excluded: list[list] = []
    shares_all = [published_share(rec, cols) for rec, _ in groups]
    total_share = sum(x for x in shares_all if x is not None) or None
    for (rec, toks), share in zip(groups, shares_all):
        cell = "" if rec["identity"] is None else str(rec["identity"]).strip()
        gcell = "" if rec["gene"] is None else str(rec["gene"]).strip()
        pct = (100.0 * share / total_share) if (share is not None and total_share) else None
        pct_txt = "" if pct is None else f"{pct:.6f}"
        entries: dict[str, dict[str, str]] = {}      # the row's entry (or entries, when a demerged token is shared)
        rank, outcome, how = "", "", ""
        summary = []
        if not toks:
            outcome = "empty_identity"
        elif toks[0][0] == "contaminant_marker":
            outcome = "contaminant_group"
        else:
            for k, (kind, acc) in enumerate(toks, 1):
                if kind != "accession":
                    summary.append(f"{k}:{kind}")
                    continue
                inf = info.get(acc)
                if inf is None:
                    summary.append(f"{k}:{acc}:not_returned")
                elif is_inactive(inf):
                    summary.append(f"{k}:{acc}:inactive")
                elif not is_reviewed_human(inf):
                    summary.append(f"{k}:{acc}:{inf['reviewed'] or '?'}/{inf['organism_id'] or '?'}")
                else:
                    summary.append(f"{k}:{acc}:reviewed_human")
                    entries, rank, how = {acc: inf}, str(k), ("primary" if k == 1 else "later_token")
                    break
            if not entries:                                                   # P3b
                for k, (kind, acc) in enumerate(toks, 1):
                    if kind == "accession" and acc in secondary:
                        entries, rank = dict(secondary[acc]), str(k)
                        how = "secondary_accession" if len(entries) == 1 else "secondary_accession_shared"
                        summary.append(f"{k}:{acc}:secondary_of:" + ";".join(sorted(entries)))
                        break
            if not entries and gene_tokens(gcell):                            # P4b, gene by gene
                per_gene: dict[str, dict[str, dict[str, str]]] = {}
                notes = []
                for g in gene_tokens(gcell):
                    hits = dict(by_gene.get(g, {}))
                    if len(hits) > 1:                                                   # M3b (D111a)
                        by_primary = {a: i for a, i in hits.items() if i.get("_matched_by") == "primary"}
                        if len(by_primary) == 1:
                            notes.append(f"{g}:{next(iter(by_primary))}(primary over synonyms {';'.join(sorted(set(hits) - set(by_primary)))})")
                            hits = by_primary
                    if len(hits) == 1:
                        per_gene[g] = hits
                        if not notes or not notes[-1].startswith(f"{g}:"):
                            notes.append(f"{g}:{next(iter(hits))}")
                    elif hits:
                        notes.append(f"{g}:ambiguous:{';'.join(sorted(hits))}")
                    else:
                        notes.append(f"{g}:unmapped")
                summary.append("gene:" + " ".join(notes))
                if per_gene and len(per_gene) == len(gene_tokens(gcell)):
                    entries = {a: {k2: v for k2, v in i.items() if k2 != "_matched_by"} for hits in per_gene.values() for a, i in hits.items()}
                    rank, how = "gene", ("gene" if len(per_gene) == 1 else "gene_shared")
            if not entries:
                outcome = "no_reviewed_human_entry"
            elif any(has_keyword(i, keratin_kw) for i in entries.values()):
                outcome = "keratin"
            elif named and any(g in named for g in gene_tokens(gcell)):
                outcome = "named_contaminant"
            elif all(a in members for a in entries):
                outcome = "member_duplicate_entry"
                for a in entries:
                    members[a]["rows"].append(str(rec["row_index"]))
            else:
                outcome = {"primary": "member", "later_token": "member_via_later_token",
                           "secondary_accession": "member_via_secondary_accession",
                           "secondary_accession_shared": "member_via_secondary_accession_shared",
                           "gene": "member_via_gene", "gene_shared": "member_via_gene_shared"}[how]
                for a, i in entries.items():
                    if a in members:
                        members[a]["rows"].append(str(rec["row_index"]))
                    else:
                        members[a] = {"info": i, "rows": [str(rec["row_index"])], "cell": cell}
        counts[outcome] = counts.get(outcome, 0) + 1
        entry_txt = ";".join(sorted(entries))
        row_table.append([rec["row_index"], cell, gcell, entry_txt, rank, outcome, "; ".join(summary), pct_txt])
        if not outcome.startswith("member") and outcome != "empty_identity":
            excluded.append([rec["row_index"], cell, gcell, entry_txt, outcome, pct_txt, pct if pct is not None else -1.0])
    excluded.sort(key=lambda r: -r[-1])
    excluded_share = sum(r[-1] for r in excluded if r[-1] >= 0)
    zero_share_members = sum(1 for row, sh in zip(row_table, shares_all) if row[5].startswith("member") and not sh)

    dat = common.category_data_dir(category)
    common.write_tsv(common.pool_file(category, pool),
                     ["accession", "gene_primary", "protein_name", "length", "reviewed", "organism_id",
                      "dataset_rows", "identity_cell"],
                     [[acc, v["info"]["gene_primary"], v["info"]["protein_name"], v["info"]["length"],
                       v["info"]["reviewed"], v["info"]["organism_id"], ";".join(v["rows"]), v["cell"]]
                      for acc, v in sorted(members.items())])
    common.write_tsv(dat / f"{pool}_dataset_rows.tsv",
                     ["row_index", "identity_cell", "gene_cell", "entry", "entry_token_rank", "outcome",
                      "tokens_checked", "published_share_percent"], row_table)
    common.write_tsv(dat / f"{pool}_rows_excluded_with_share.tsv",
                     ["row_index", "identity_cell", "gene_cell", "entry", "outcome", "published_share_percent"],
                     [r[:-1] for r in excluded])
    log.info("pool %s: %d rows -> %s; %d entries; excluded share %.4f %% of the listing", pool, len(rows), counts, len(members), excluded_share)
    record = {"dataset_source": src, "dataset_file": fkey, "dataset_path": m["path"], "dataset_sha256": got,
              "sheet": cols["sheet"], "header_row": cols["header_row"], "identity_column": cols["identity_column"],
              "gene_column": cols["gene_column"], "keratin_keyword_id": keratin_kw,
              "named_contaminant_genes": ", ".join(sorted(named)) or "(none)",
              "rows": str(len(rows)), "accessions_looked_up": str(len(info)),
              "tokens_inactive_in_uniprot": str(sum(1 for i in info.values() if is_inactive(i))),
              "tokens_asked_as_secondary": str(len(unresolved)), "tokens_resolved_as_secondary": str(len(secondary)),
              "genes_asked_as_fallback": str(len(fallback_genes)), "pool_size": str(len(members)),
              "excluded_share_percent_of_listing": f"{excluded_share:.6f}",
              "member_rows_with_zero_or_no_share": str(zero_share_members),
              **{f"n_{k}": str(v) for k, v in sorted(counts.items())},
              "pool_file": common.pool_file(category, pool).name, "rows_file": f"{pool}_dataset_rows.tsv",
              "excluded_file": f"{pool}_rows_excluded_with_share.tsv"}
    return set(members), record


def main_category(category: str, log) -> int:
    files = common.category_files(category)
    decisions = common.read_ini(files["decisions"])
    if not decisions.has_section("contaminant_rules") or not decisions["contaminant_rules"].get("keratin_keyword_id", "").strip():
        raise SystemExit(f"[STOP] {files['decisions']} needs [contaminant_rules] keratin_keyword_id (D98, D107)")
    rules = decisions["contaminant_rules"]
    pools = common.pool_names(category)
    client = UniProtClient(log)
    queries = common.new_ini()
    queries["run"] = {"date": common.iso_now(), "category": category,
                      "query_template": "(accession:A1 OR accession:A2 ... )", "fields": ",".join(LOOKUP_FIELDS),
                      "endpoint": "https://rest.uniprot.org/uniprotkb/search"}
    sets: dict[str, set[str]] = {}
    for pool in pools:
        sets[pool], queries[f"pool.{pool}"] = enumerate_published_pool(client, category, pool, decisions[f"pool.{pool}"], rules, log)
    queries["run"]["uniprot_release"] = client.release or "(header not returned)"
    common.write_ini(queries, files["pool_queries"], [
        "pool_queries.ini -- GENERATED by enumerate_pool.py. DO NOT EDIT BY HAND.",
        f"Record of the UniProt lookups that produced data/{category}/pool_<pool>.tsv (rules P1-P6, D107).",
    ])
    overlap_rows = []
    names = sorted(sets)
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            for acc in sorted(sets[a] & sets[b]):
                overlap_rows.append([acc, a, b])
    common.write_tsv(files["protein_set_out"] / "pool_overlap.tsv", ["accession", "pool_a", "pool_b"], overlap_rows)
    log.info("pool overlap: %d accessions in more than one pool", len(overlap_rows))
    log.info("UniProt release: %s", client.release)
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--category", default=None, help="a non-muscle category: enumerate its [pool.<name>] pools by published accession (P1-P6)")
    args = ap.parse_args(argv)
    log = common.make_logger(STAGE if args.category is None else f"{STAGE}_{args.category}")
    if args.category:
        return main_category(args.category, log)
    reset_flags(common.FLAGS_TSV, STAGE)
    resolved = common.read_ini(common.RESOLVED_TERMS_INI)
    client = UniProtClient(log)

    queries = common.new_ini()
    queries["run"] = {"date": common.iso_now(), "query_template": term_query("GO:NNNNNNN"),
                      "fields": ",".join(FIELDS), "endpoint": "https://rest.uniprot.org/uniprotkb/search"}
    pools: dict[str, set[str]] = {}
    for section in resolved.sections():
        if not section.startswith("tier."):
            continue
        tier = section.split(".", 1)[1]
        if resolved[section].get("definition") == "measured_remainder":
            continue                                    # after every ontology tier (M4)
        term_id = resolved[section]["term_id"]
        term_name = resolved[section]["term_name"]
        subtree = common.read_tsv(common.ONTOLOGY_DIR / f"tier{tier}_subtree.tsv")
        neighbours = common.read_tsv(common.ONTOLOGY_DIR / f"tier{tier}_neighbours.tsv")
        pool, n_top, ok = enumerate_tier(client, tier, term_id, term_name, subtree, neighbours, log)
        pools[tier] = pool
        queries[section] = {"term_id": term_id, "term_name": term_name,
                            "n_subtree_terms": str(len(subtree)), "pool_size": str(len(pool)),
                            "term_level_query_size": str(n_top),
                            "descendant_expansion_matches_union": "true" if ok else "false",
                            "pool_file": f"tier{tier}_pool.tsv"}
    for section in resolved.sections():
        if section.startswith("tier.") and resolved[section].get("definition") == "measured_remainder":
            tier = section.split(".", 1)[1]
            pool, counts = enumerate_measured_tier(client, tier, resolved[section], pools, log)
            pools[tier] = pool
            queries[section] = {"definition": "measured_remainder", "dataset_source": resolved[section]["dataset_source"],
                                "dataset_file": resolved[section]["dataset_file"],
                                "contaminant_source": resolved[section]["contaminant_source"],
                                "query_template": "(organism_id:9606) AND (reviewed:true) AND (gene_exact:G1 OR ... )",
                                "genes": str(sum(counts.values())), "pool_size": str(len(pool)),
                                **{f"n_{k}": str(v) for k, v in counts.items()},
                                "pool_file": f"tier{tier}_pool.tsv", "mapping_file": f"tier{tier}_gene_mapping.tsv"}
    queries["run"]["uniprot_release"] = client.release or "(header not returned)"
    common.write_ini(queries, common.POOL_QUERIES_INI, [
        "pool_queries.ini -- GENERATED by enumerate_pool.py. DO NOT EDIT BY HAND.",
        "Record of the UniProt queries that produced data/tier<N>_pool.tsv.",
    ])

    # overlap between tiers
    overlap_rows = []
    tiers = sorted(pools)
    for i, a in enumerate(tiers):
        for b in tiers[i + 1:]:
            for acc in sorted(pools[a] & pools[b]):
                overlap_rows.append([acc, a, b])
    common.write_tsv(common.protein_set_out_dir() / "pool_overlap.tsv", ["accession", "tier_a", "tier_b"], overlap_rows)
    log.info("tier overlap: %d accessions in more than one tier", len(overlap_rows))
    log.info("UniProt release: %s", client.release)
    return 0


if __name__ == "__main__":
    sys.exit(main())
