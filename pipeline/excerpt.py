#!/usr/bin/env python3
"""
excerpt.py — tooling, not a pipeline stage: bounded excerpts of the large generated files.

The generated tables run to thousands of rows; a review in chat needs their headers and a
few hundred well-chosen rows, not the bodies. This tool writes, for a fixed list of files,
either the whole file (small ones) or a labelled cut of it — the top N rows by a column,
the first N rows of an already-ranked table, the rows matching a filter, or the rows whose
key appears in an earlier cut — keeping every `#` header line so the provenance travels
with the rows. Nothing is computed, changed, or interpreted; every excerpt states which
file it came from, the rule that selected its rows, and how many of the file's rows it holds.

    python run.py excerpt

Writes  excerpts/<UTC stamp>/<same relative path>[__<cut>].<ext>
        excerpts/<UTC stamp>/INDEX.md                 what each excerpt is
        excerpts/excerpts_<UTC stamp>.tar.gz          the folder, for upload
        logs/excerpt_<UTC>.log

Source files are the generated TSV/INI tables and, since Step 7b, the literature workbooks
themselves (.xlsx, .xls). A workbook cut names its sheet and which rows form the header;
the excerpt is written as TSV with the rows above the header reproduced verbatim and a
letter-to-column map, so a column map in config can be written against what the file holds
rather than against a note about it. Nothing is computed, changed, or interpreted.

`excerpts/` is not committed and is not part of the record (like docs/handoffs/). A file
that is missing on this machine is listed as missing in INDEX.md and the log; it does not
stop the run. The list of files and cuts is SPEC below — edit it there.
"""

from __future__ import annotations

import sys
import tarfile
from pathlib import Path

from pipeline import common

EXCERPTS_DIR = common.REPO_ROOT / "excerpts"

# ----------------------------------------------------------------------------
# What to excerpt. Each entry: (relative path, [cuts]). A cut is a dict:
#   {"kind": "whole"}
#   {"kind": "head",   "n": N}                                  first N rows of an already-ranked table
#   {"kind": "top",    "n": N, "by": <column>}                  N largest by a numeric column
#   {"kind": "top",    "n": N, "by_col": "P"}                   the same, by spreadsheet column letter
#   {"kind": "equals", "column": c, "value": v}                 rows where column == value
#   {"kind": "contains", "column": c, "value": v}               rows where value is a substring of the cell
#   {"kind": "in", "column": c, "from": <path>, "from_cut": <name>, "from_column": c2}
#                                                              rows whose column is in the set of
#                                                              values of c2 in an earlier cut
#   {"kind": "sheet"}                                           a whole sheet of a workbook, as TSV
#   {"kind": "sheet_list"}                                      every sheet of a workbook with its size
# Every cut has a "name" used in the output filename; "whole" needs none.
#
# A cut on a workbook (.xlsx, .xls) also carries:
#   "sheet": <sheet name>          which sheet — a name that is not in the workbook is reported
#                                  with the names that are, and does not stop the run
#   "header_rows": [2, 3]          1-based sheet rows that form the header (default [1]). Rows
#                                  before the last one are forward-filled so a merged group label
#                                  reaches the columns it spans; names are joined with ' | '.
# {"optional": True} on any cut: a file that is a conditional output of its stage, so its
# absence is reported without a warning.
# ----------------------------------------------------------------------------

MF = "outputs/intermediate/mass_fractions"
CO = "outputs/intermediate/composition"

# ----------------------------------------------------------------------------
# One spec per command, in run order. A stage's excerpt list sits with the command that
# writes those files, so adding an output means adding its cut in one obvious place rather
# than hunting through a single flat list. `python run.py excerpt` concatenates them.
# ----------------------------------------------------------------------------

SPECS: dict[str, list[tuple[str, list[dict]]]] = {
    "fetch-literature": [
        ("data/literature/manifest.ini", [{"kind": "whole"}]),
        # Step 7b, the blood sources: read before any blood stage code, so the column maps in
        # config/blood/mass_fraction_decisions.ini are written against what the files hold.
        ("data/literature/bryk_2017/pr7b00025_si_002.xlsx", [
            {"kind": "sheet_list", "name": "sheets"},
            {"kind": "head", "name": "first20", "n": 20, "sheet": "Table S3", "header_rows": [2, 3]},
            {"kind": "top", "name": "top100_by_col_P", "n": 100, "by_col": "P",
             "sheet": "Table S3", "header_rows": [2, 3]},
        ]),
        ("data/literature/geyer_2016/1-s2.0-S2405471216300722-mmc3.xlsx", [
            {"kind": "sheet_list", "name": "sheets"},
            {"kind": "sheet", "name": "table_s2", "sheet": "Supplementary Table S2"},
        ]),
        ("data/literature/geyer_2016/1-s2.0-S2405471216300722-mmc6.xlsx", [
            {"kind": "sheet_list", "name": "sheets"},
            {"kind": "sheet", "name": "table_s5"},
        ]),
        ("data/literature/hortin_2008/clinchem.2008.108175-2.xls", [
            {"kind": "sheet_list", "name": "sheets"},
            {"kind": "sheet", "name": "database", "sheet": "DataBase"},
        ]),
    ],
    "protein-set": [
        ("data/gene-ontology/source.ini", [{"kind": "whole"}]),
        ("data/iupac/source.ini", [{"kind": "whole"}]),
        ("data/iupac/amino_acid_symbols.ini", [{"kind": "whole"}]),
        ("data/pubchem/amino_acid_masses.ini", [{"kind": "whole"}]),
        ("outputs/flags.tsv", [{"kind": "whole"}]),
        ("config/mass_fractions_per_entry.tsv", [
        {"kind": "top", "name": "top100_by_w_I_combined", "n": 100, "by": "w_I_combined"},
        {"kind": "equals", "name": "no_dataset_row", "column": "match_rule", "value": "none"},
        {"kind": "contains", "name": "shared_split", "column": "match_rule", "value": "shared_split"},
        {"kind": "contains", "name": "via_mapping", "column": "match_rule", "value": "via_mapping"},
        ]),
        ("outputs/intermediate/protein_set/isoform_bound.tsv", [{"kind": "whole"}]),
        ("outputs/intermediate/protein_set/processing_bound.tsv", [{"kind": "whole"}]),
        ("outputs/intermediate/protein_set/excluded_non_standard_alphabet.tsv", [{"kind": "whole"}]),
        ("outputs/intermediate/protein_set/chain_positions_resolved.tsv", [{"kind": "whole"}]),
    ],
    "composition": [
        (f"{CO}/composition_summary.ini", [{"kind": "whole"}]),
        (f"{CO}/ptm_summary.ini", [{"kind": "whole"}]),
        (f"{CO}/amino_acid_composition_per_protein.tsv", [
        {"kind": "in", "name": "master_rows_of_top100", "column": "accession",
        "from": "config/mass_fractions_per_entry.tsv", "from_cut": "top100_by_w_I_combined", "from_column": "accession",
        "also_equals": ("segment_set", "master")},
        ]),
        (f"{CO}/processing_mass_bound.tsv", [{"kind": "whole"}]),
    ],
    "mass-fractions": [
        (f"{MF}/mass_fractions_summary.ini", [{"kind": "whole"}]),
        ("outputs/intermediate/digest/digest_summary.ini", [{"kind": "whole"}]),
        ("outputs/intermediate/literature_inventory/literature_inventory_summary.ini", [{"kind": "whole"}]),
        (f"{MF}/combined_entries_ranked.tsv", [{"kind": "head", "name": "first50", "n": 50}]),
        (f"{MF}/tier1_entries_ranked.tsv", [{"kind": "head", "name": "first30", "n": 30}]),
        (f"{MF}/tier2_entries_ranked.tsv", [{"kind": "head", "name": "first30", "n": 30}]),
        (f"{MF}/dataset_rows_outside_pool.tsv", [{"kind": "head", "name": "first30", "n": 30}]),
        (f"{MF}/weighted_bounds.tsv", [{"kind": "whole"}]),
        (f"{MF}/family_bounds.tsv", [{"kind": "whole"}]),
        (f"{MF}/band_families.tsv", [{"kind": "whole"}]),
        (f"{MF}/classical_check_carroll_2004.tsv", [{"kind": "whole"}]),
        (f"{MF}/excluded_entries_mass_share.tsv", [{"kind": "whole"}]),
        (f"{MF}/branch_exclusive_mass.tsv", [{"kind": "whole"}]),
        (f"{MF}/pool_entries_without_dataset_row.tsv", [{"kind": "whole"}]),
        (f"{MF}/momenzadeh_2023_myh_fractions_ibaq_vs_lfq.tsv", [{"kind": "whole"}]),
        (f"{MF}/ratio_check_moreno-justicia_2025.tsv", [{"kind": "head", "name": "first20", "n": 20}]),
        (f"{MF}/ratio_check_deshmukh_2021.tsv", [{"kind": "head", "name": "first20", "n": 20}]),
        ("outputs/intermediate/digest/density_ranked.tsv", [{"kind": "head", "name": "first20", "n": 20}]),
        ("outputs/intermediate/digest/families.tsv", [{"kind": "head", "name": "first40", "n": 40}]),
    ],
    "standard": [
        ("outputs/standard/_calculated_amino_acid_standard.tsv", [{"kind": "whole"}]),
        ("outputs/standard/_calculated_amino_acid_standard_residue_convention.tsv", [{"kind": "whole"}]),
        ("outputs/standard/standard_summary.ini", [{"kind": "whole"}]),
        ("outputs/standard/_calculated_amino_acid_standard_with_non_protein_metabolite_pools.tsv", [{"kind": "whole"}]),
        ("outputs/standard/non_protein_metabolite_pools_summary.ini", [{"kind": "whole"}]),
        ("outputs/standard/non_protein_metabolite_pool_adjustment_per_amino_acid.tsv", [{"kind": "whole"}]),
        ("outputs/standard/non_protein_metabolite_pool_amounts_per_kg_muscle.tsv", [{"kind": "whole"}]),
        ("outputs/standard/sensitivity_non_protein_metabolite_pool_turnover_frame.tsv", [{"kind": "whole", "optional": True}]),
        ("outputs/standard/sensitivity_non_protein_metabolite_pool_basis.tsv", [{"kind": "whole", "optional": True}]),
        ("outputs/standard/sensitivity_non_protein_metabolite_pool_sex.tsv", [{"kind": "whole", "optional": True}]),
        ("outputs/standard/sensitivity/sensitivity_non_protein_metabolite_pool_spread.tsv", [{"kind": "whole"}]),
        ("outputs/standard/eaa_subset_with_non_protein_metabolite_pools.tsv", [{"kind": "whole"}]),
        ("outputs/standard/stress/stress_summary.ini", [{"kind": "whole"}]),
        ("outputs/standard/stress/stress_summary_per_scenario.tsv", [{"kind": "whole"}]),
        ("outputs/standard/stress/stress_influence_per_entry.tsv", [{"kind": "head", "name": "first60", "n": 60}]),
        ("outputs/standard/stress/composition_distance_top_entries.tsv", [{"kind": "whole"}]),
        ("outputs/standard/uncertainty/uncertainty_summary.ini", [{"kind": "whole"}]),
        ("outputs/standard/amino_acid_profiles.tsv", [{"kind": "whole"}]),
        ("outputs/standard/amino_acid_g_per_100g_protein_free.tsv", [{"kind": "whole"}]),
        ("outputs/standard/amino_acid_g_per_100g_protein_residue.tsv", [{"kind": "whole"}]),
        ("outputs/standard/profile_differences.tsv", [{"kind": "whole"}]),
        ("outputs/standard/sensitivity/sensitivity_mhc_actin_profiles.tsv", [{"kind": "whole"}]),
        ("outputs/standard/sensitivity/sensitivity_mhc_actin_spread.tsv", [{"kind": "whole"}]),
        ("outputs/standard/uncertainty/bounds_after_weighting.tsv", [{"kind": "whole"}]),
        ("outputs/standard/sensitivity/completeness_sensitivity.tsv", [{"kind": "whole"}]),
        ("outputs/standard/eaa_subset.tsv", [{"kind": "whole"}]),
        ("outputs/standard/uncertainty/uncertainty_intervals.tsv", [{"kind": "whole"}]),
        ("outputs/standard/uncertainty/uncertainty_per_amino_acid.tsv", [{"kind": "whole"}]),
        ("outputs/standard/uncertainty/sd_to_median_ratio.tsv", [
        {"kind": "top", "name": "top100_by_w_combined", "n": 100, "by": "w_combined"},
        {"kind": "top", "name": "top100_by_sd_over_median", "n": 100, "by": "sd_over_median"},
        ]),
    ],
    "usda": [
        ("outputs/usda/usda_archive_summary.ini", [{"kind": "whole"}]),
        ("outputs/usda/usda_nutrient_map.tsv", [{"kind": "whole"}]),
        ("outputs/usda/amino_acids_per_food.tsv", [{"kind": "head", "name": "first60", "n": 60}]),
        ("outputs/usda/foods_without_amino_acids.tsv", [{"kind": "head", "name": "first40", "n": 40}]),
    ],
    "comparison": [
        ("outputs/comparison/gorissen_2018_human_muscle.tsv", [{"kind": "whole"}]),
        ("outputs/comparison/calculated_vs_gorissen_2018.tsv", [{"kind": "whole"}]),
        ("outputs/comparison/gorissen_2018_mass_balance.tsv", [{"kind": "whole"}]),
        ("outputs/comparison/comparison_summary.ini", [{"kind": "whole"}]),
    ],
    "match": [
        ("outputs/match/match_summary.ini", [{"kind": "whole"}]),
        ("outputs/match/limiting_amino_acid_counts.tsv", [{"kind": "whole"}]),
        ("outputs/match/match_rate_by_reference.tsv", [{"kind": "head", "name": "first60", "n": 60}]),
        ("outputs/match/match_rate_per_food.tsv", [{"kind": "head", "name": "first60", "n": 60}]),
        ("outputs/match/match_rate_steps.tsv", [{"kind": "head", "name": "first60", "n": 60}]),
        ("outputs/match/match_rate_ranked_skeletal_muscle_protein_with_non_protein_metabolite_pools.tsv", [{"kind": "head", "name": "first100", "n": 100}]),
        ("outputs/match/match_rate_ranked_gorissen_2018_human_muscle.tsv", [{"kind": "head", "name": "first100", "n": 100}]),
        ("outputs/match/foods_not_scored.tsv", [{"kind": "head", "name": "first40", "n": 40}]),
    ],
}

# The flat list the runner walks: every command's spec, in run order.
SPEC: list[tuple[str, list[dict]]] = [entry for cmd in SPECS for entry in SPECS[cmd]]


# ----------------------------------------------------------------------------

def read_table(path: Path) -> tuple[list[str], list[str], list[list[str]]]:
    """(comment lines, header columns, rows) of a TSV whose first lines may be '#' comments."""
    comments, header, rows = [], None, []
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.rstrip("\n")
            if header is None and line.startswith("#"):
                comments.append(line)
                continue
            if not line:
                continue
            if header is None:
                header = line.split("\t")
                continue
            rows.append(line.split("\t"))
    return comments, header or [], rows


def as_float(x: str) -> float:
    try:
        return float(x)
    except ValueError:
        return float("-inf")


# ---------------------------------------------------------------- workbooks (Step 7b)

WORKBOOK_EXT = {".xlsx", ".xls"}


class SheetNotFound(Exception):
    """The cut names a sheet the workbook does not have; the message lists the ones it does."""


def col_letter(i: int) -> str:
    """1 -> A, 26 -> Z, 27 -> AA: the column letter a person reads in the spreadsheet."""
    out = ""
    while i > 0:
        i, r = divmod(i - 1, 26)
        out = chr(65 + r) + out
    return out


def col_index(letter: str) -> int:
    """A -> 0, P -> 15: the 0-based position of a spreadsheet column letter."""
    n = 0
    for ch in letter.strip().upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def cell_text(v) -> str:
    """A cell as text. A whole-number float is written without its '.0' — an accession count
    or a row index read back as 153.0 is noise; nothing else is rounded or reformatted."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return str(v).replace("\t", " ").replace("\n", " ").replace("\r", " ").strip()


def trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop trailing empty cells and trailing empty rows: a sheet's used range often
    reports further than its content, and the padding would travel into the excerpt."""
    out = []
    for row in rows:
        while row and row[-1] == "":
            row = row[:-1]
        out.append(row)
    while out and not any(out[-1]):
        out.pop()
    return out


def workbook_sheets(path: Path) -> list[tuple[str, int, int]]:
    """(sheet name, rows, columns) for every sheet, in the workbook's own order.

    The two numbers are the sheet's USED RANGE as the file reports it, which is what a
    workbook can answer without reading every cell. It can exceed the content — a sheet
    whose rows were formatted once reports them for ever. The row count to trust is the one
    in an excerpt of that sheet, which counts rows that hold something.
    """
    if path.suffix.lower() == ".xls":
        import xlrd                                    # .xls only; openpyxl cannot read it
        book = xlrd.open_workbook(path)
        return [(s.name, s.nrows, s.ncols) for s in book.sheets()]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, ws.max_row or 0, ws.max_column or 0) for ws in wb.worksheets]
    finally:
        wb.close()


def read_sheet(path: Path, sheet: str | None) -> tuple[str, list[list[str]]]:
    """(sheet name, every row as text). The named sheet, or the first when none is named."""
    names = [n for n, _, _ in workbook_sheets(path)]
    if not names:
        raise SheetNotFound(f"{path.name} has no sheets")
    if sheet is None:
        name = names[0]
    else:
        match = [n for n in names if n.strip().lower() == sheet.strip().lower()]
        if not match:
            raise SheetNotFound(f"{path.name} has no sheet {sheet!r}; its sheets are: " + ", ".join(repr(n) for n in names))
        name = match[0]
    if path.suffix.lower() == ".xls":
        import xlrd
        sh = xlrd.open_workbook(path).sheet_by_name(name)
        rows = [[cell_text(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = [[cell_text(v) for v in row] for row in wb[name].iter_rows(values_only=True)]
        finally:
            wb.close()
    return name, trim(rows)


def forward_fill(row: list[str]) -> list[str]:
    """Carry each value into the empty cells that follow it: a merged group label is stored
    in its first cell only, and the columns it spans would otherwise look unlabelled."""
    out, last = [], ""
    for v in row:
        last = v or last
        out.append(last)
    return out


def sheet_header(rows: list[list[str]], header_rows: list[int]) -> list[str]:
    """Column names from the 1-based sheet rows that form the header. Every row but the last
    is padded to the sheet's full width and forward-filled (group labels); the last is taken as
    it is, so a genuinely empty column stays empty and is named by its letter.

    The padding matters: a group-label row is usually shorter than the row of names under it —
    trailing empty cells are not stored — and filling only to that row's own end leaves the
    columns past it unlabelled. In bryk_2017's Table S3 that silently dropped 'White ghosts'
    from the last three copy-number columns.
    """
    width = max((len(r) for r in rows), default=0)
    picked = []
    for n, r in enumerate(header_rows, 1):
        row = list(rows[r - 1]) if r - 1 < len(rows) else []
        row += [""] * (width - len(row))
        picked.append(forward_fill(row) if n < len(header_rows) else row)
    names = []
    for j in range(width):
        parts = []
        for r in picked:
            v = r[j] if j < len(r) else ""
            if v and (not parts or parts[-1] != v):
                parts.append(v)
        names.append(" | ".join(parts) or col_letter(j + 1))
    return names


def sheet_table(path: Path, cut: dict) -> tuple[list[str], list[str], list[list[str]], str]:
    """(comment lines, header, body rows, sheet name) for a workbook cut. The comment lines
    reproduce every sheet row above the header verbatim and map each column letter to its
    name, so the excerpt says what the spreadsheet says."""
    name, rows = read_sheet(path, cut.get("sheet"))
    header_rows = list(cut.get("header_rows", [1]))
    header = sheet_header(rows, header_rows)
    last = max(header_rows)
    comments = [f"# sheet {name!r}; header row(s) {', '.join(str(r) for r in header_rows)}; "
                f"{len(rows) - last} data rows below row {last}"]
    for i in range(1, last):
        comments.append(f"# sheet row {i}: " + " | ".join(rows[i - 1]) if i - 1 < len(rows) else f"# sheet row {i}: (empty)")
    comments.append("# columns: " + "; ".join(f"{col_letter(j + 1)} = {n}" for j, n in enumerate(header)))
    body = [r + [""] * (len(header) - len(r)) for r in rows[last:]]
    return comments, header, body, name


def apply_cut(cut: dict, header: list[str], rows: list[list[str]], collected: dict) -> tuple[list[list[str]], str]:
    """Return (rows kept, one-line description of the rule)."""
    kind = cut["kind"]
    if kind == "head":
        return rows[: cut["n"]], f"first {cut['n']} rows as ordered in the file"
    if kind == "sheet":
        return rows, "the whole sheet"
    if kind == "top":
        if "by_col" in cut:
            i = col_index(cut["by_col"])
            label = f"column {cut['by_col']} ({header[i] if i < len(header) else '?'})"
        else:
            i = header.index(cut["by"])
            label = f"column {cut['by']}"
        return sorted(rows, key=lambda r: -as_float(r[i] if i < len(r) else ""))[: cut["n"]], \
            f"the {cut['n']} largest by {label} (descending)"
    if kind == "equals":
        i = header.index(cut["column"])
        return [r for r in rows if r[i] == cut["value"]], f"rows where {cut['column']} == {cut['value']!r}"
    if kind == "contains":
        i = header.index(cut["column"])
        return [r for r in rows if cut["value"] in r[i]], f"rows where {cut['column']} contains {cut['value']!r}"
    if kind == "in":
        keys = collected.get((cut["from"], cut["from_cut"], cut["from_column"]))
        if keys is None:
            raise KeyError(f"cut {cut['name']}: its source cut {cut['from']}::{cut['from_cut']} did not run first")
        i = header.index(cut["column"])
        kept = [r for r in rows if r[i] in keys]
        rule = f"rows whose {cut['column']} is one of the {len(keys)} values of {cut['from_column']} in excerpt {cut['from']}::{cut['from_cut']}"
        if "also_equals" in cut:
            c, v = cut["also_equals"]
            j = header.index(c)
            kept = [r for r in kept if r[j] == v]
            rule += f", and {c} == {v!r}"
        return kept, rule
    raise ValueError(f"unknown cut kind {kind!r}")


def main(argv: list[str] | None = None) -> int:
    log = common.make_logger("excerpt")
    stamp = common.timestamp()
    out_root = EXCERPTS_DIR / stamp
    out_root.mkdir(parents=True, exist_ok=True)
    index = ["# Excerpts", "",
             f"Written {common.iso_now()} by `pipeline/excerpt.py` from the files on this machine. "
             "Each excerpt keeps its source file's `#` header lines and adds one line naming the rule that "
             "selected its rows. Nothing is computed or changed here.", "",
             "| Excerpt | Source file | Rule | Rows kept / in file |", "|---|---|---|---|"]
    collected: dict[tuple, set[str]] = {}
    written = 0
    for rel, cuts in SPEC:
        src = common.REPO_ROOT / rel
        if not src.exists():
            if all(c.get("optional") for c in cuts):
                log.info("not written by this run (conditional output): %s", rel)
                index.append(f"| — | `{rel}` | not written by this run (conditional output) | — |")
            else:
                log.warning("missing on this machine: %s", rel)
                index.append(f"| — | `{rel}` | missing on this machine | — |")
            continue
        is_workbook = src.suffix.lower() in WORKBOOK_EXT
        for cut in cuts:
            if is_workbook and cut["kind"] == "sheet_list":
                sheets = workbook_sheets(src)
                p = Path(rel)
                dest = out_root / p.parent / f"{p.stem}__{cut['name']}.tsv"
                dest.parent.mkdir(parents=True, exist_ok=True)
                lines = [f"# EXCERPT of {rel}: every sheet with its size; written {common.iso_now()} by pipeline/excerpt.py",
                         "# rows/columns are the sheet's used range as the file reports it; it can exceed the "
                         "content. The count that holds is the one in an excerpt of the sheet.",
                         "sheet\trows_used_range\tcolumns_used_range"] + [f"{n}\t{r}\t{c}" for n, r, c in sheets]
                dest.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
                index.append(f"| `{dest.relative_to(out_root).as_posix()}` | `{rel}` | every sheet with its used range | {len(sheets)} sheets |")
                log.info("sheets %s -> %d sheet(s): %s", rel, len(sheets), ", ".join(n for n, _, _ in sheets))
                written += 1
                continue
            if is_workbook:
                try:
                    comments, header, rows, sheet_name = sheet_table(src, cut)
                except SheetNotFound as e:
                    log.warning("%s: %s", rel, e)
                    index.append(f"| — | `{rel}` | {e} | — |")
                    continue
                kept, rule = apply_cut(cut, header, rows, collected)
                rule = f"sheet {sheet_name!r}, {rule}"
                p = Path(rel)
                dest = out_root / p.parent / f"{p.stem}__{cut['name']}.tsv"
                dest.parent.mkdir(parents=True, exist_ok=True)
                lines = [f"# EXCERPT of {rel}: {rule}; {len(kept)} of {len(rows)} rows; written {common.iso_now()} by pipeline/excerpt.py"]
                lines += comments + ["\t".join(header)] + ["\t".join(r) for r in kept]
                dest.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
                index.append(f"| `{dest.relative_to(out_root).as_posix()}` | `{rel}` | {rule} | {len(kept)} / {len(rows)} |")
                log.info("cut    %s :: %s -> %d of %d rows", rel, cut["name"], len(kept), len(rows))
                written += 1
                continue
            if cut["kind"] == "whole":
                dest = out_root / rel
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes(src.read_bytes())
                n = sum(1 for _ in open(src, encoding="utf-8", errors="replace"))
                index.append(f"| `{rel}` | `{rel}` | whole file | {n} lines |")
                log.info("whole  %s (%d lines)", rel, n)
                written += 1
                continue
            comments, header, rows = read_table(src)
            kept, rule = apply_cut(cut, header, rows, collected)
            name = cut["name"]
            p = Path(rel)
            dest = out_root / p.parent / f"{p.stem}__{name}{p.suffix}"
            dest.parent.mkdir(parents=True, exist_ok=True)
            lines = [f"# EXCERPT of {rel}: {rule}; {len(kept)} of {len(rows)} rows; written {common.iso_now()} by pipeline/excerpt.py"]
            lines += comments + ["\t".join(header)] + ["\t".join(r) for r in kept]
            dest.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")
            for col in header:                       # every column of every cut is collectable by later cuts
                i = header.index(col)
                collected[(rel, name, col)] = {r[i] for r in kept}
            index.append(f"| `{dest.relative_to(out_root).as_posix()}` | `{rel}` | {rule} | {len(kept)} / {len(rows)} |")
            log.info("cut    %s :: %s -> %d of %d rows", rel, name, len(kept), len(rows))
            written += 1
    (out_root / "INDEX.md").write_text("\n".join(index) + "\n", encoding="utf-8", newline="\n")
    tar_path = EXCERPTS_DIR / f"excerpts_{stamp}.tar.gz"
    with tarfile.open(tar_path, "w:gz") as tf:
        tf.add(out_root, arcname=stamp)
    log.info("%d excerpts in %s; tarball %s (%d bytes)", written, out_root.relative_to(common.REPO_ROOT).as_posix(),
             tar_path.relative_to(common.REPO_ROOT).as_posix(), tar_path.stat().st_size)
    print(f"\nUPLOAD: {tar_path.relative_to(common.REPO_ROOT).as_posix()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())