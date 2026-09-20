#!/usr/bin/env python3
"""
score.py — tooling: Match Rate for a workbook of foods entered by hand.

NOT a pipeline stage. It is not in `common.COMMAND_ORDER`, nothing downstream reads what it
writes, and no currency test rebuilds it. It is the ad-hoc path: transcribe a label, paste a
supplier's analysis, enter a finished blend, and get the Match Rate back without editing config,
without rerunning `match` over every USDA food, and without any of it entering the repository's
record. `scoring/` is gitignored for that reason (M7).

    1. put a workbook in  scoring/
    2. python run.py score
    3. read  scoring/<name>_scored.xlsx

The arithmetic is `match.match_rate()` — the same function `pipeline/match.py` scores USDA with
(M6). Nothing about the formula is reimplemented here; what is here is the workbook in, the
workbook out, and the imputation rule for an amino acid a label does not print (M8).

The reference is whatever `config/match_rate.ini [formula] summary_reference` names — the most
complete standard built so far (D128). As categories 3 through 11 land, that key moves and this
tool follows it; no reference is named here or in the workbook.

Reads
  scoring/*.xlsx                                the workbooks you put there (`_scored` files skipped)
  config/match_rate.ini                         the reference, its scored set, the version label
  config/fao_2013_indispensable_amino_acids.ini the nine headings (D62), reduced per D86
  data/iupac/amino_acid_symbols.ini             three-letter / trivial name -> letter: nothing is named in code
  the standard table the reference names         its values

Writes (scoring/, beside the input)
  <name>_scored.xlsx   sheet `match_rate`: YOUR rows, every column and value exactly as entered and in
                                           your order, with the calculation columns appended to the right
                       sheet `steps`:      the spreadsheet's own walk per product (M5), so any score can
                                           be rebuilt by hand
                       sheet `reference`:  what was scored against, where it was read, its hash, and its
                                           values — the provenance, as rows and columns, not as comments

The input workbook is never written to.

Columns of the input sheet
  label         required. The product's identity. Duplicates are allowed: two rows, two scores.
  basis         required. `product` or `protein` — what the amounts are per.
  grams_basis   required. How many grams of that basis the amounts belong to (100, 250, 31.5 ...).
  protein_g     optional. Protein in that same amount. Blank leaves the density columns blank.
  <Xxx>         the amino acids, three-letter symbol or trivial name (His, Ile, Trp / Histidine ...),
                in the same unit as grams_basis. BLANK means not reported and is imputed (M8);
                0 means measured as zero and scores zero (M3).
  anything else carried through to the output untouched: source, location, note, a supplier's code.

Rules M7-M10 in docs/conventions.md.

Usage
  python run.py score          # the whole scoring/ folder
  python pipeline/score.py     # the same
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pipeline import common
from pipeline.match import (Stop, load_references, load_scored_sets, match_rate,
                            read_config_reference, read_standard_column, sheet_steps, symbol_tables)
from pipeline.mass_fractions import sha256_path

AA = list(common.AMINO_ACIDS)
CONFIG_INI = common.MATCH_RATE_INI
SCORED_SUFFIX = "_scored"
GENERATED_BY = "score.py"

REQUIRED_COLUMNS = ["label", "basis", "grams_basis"]
BASES = ("product", "protein")

# The calculation columns, appended to the right of the author's own. Order is the reading order:
# the score first, then what it was scored on, then the totals, then the ratio per amino acid.
CALC_COLUMNS = ["match_rate_percent", "limiting_amino_acid", "scored_amino_acids",
                "imputed_at_reference_share", "not_scored_reason",
                "eaa_total", "eaa_percent_of_protein", "eaa_g_per_100g_product",
                "protein_percent_of_product"]


def _rel(p: Path) -> str:
    try:
        return p.relative_to(common.REPO_ROOT).as_posix()
    except ValueError:
        return p.name


def _path(rel: str) -> Path:
    p = Path(rel.strip())
    return p if p.is_absolute() else common.REPO_ROOT / p


# --------------------------------------------------------------------------- the reference (one, D128)

def load_summary_reference() -> dict:
    """The reference `[formula] summary_reference` names, with its scored set resolved.

    The same loaders match.py uses, so a reference that scores in the pipeline scores here and a
    rename in config reaches both at once.
    """
    if not CONFIG_INI.exists():
        raise Stop(f"{_rel(CONFIG_INI)} not found")
    cp = common.read_ini(CONFIG_INI)
    try:
        three, names = symbol_tables()
    except FileNotFoundError:
        raise Stop(f"{_rel(common.AMINO_ACID_SYMBOLS_INI)} is not on this machine — it is the IUPAC-IUBMB "
                   f"table every stage resolves amino acid names through. Run `python run.py composition` "
                   f"first (docs/run_order.md).")
    scored_sets = load_scored_sets(cp, three, names)
    refs = load_references(cp, scored_sets, three, names)
    name = cp["formula"].get("summary_reference", "").strip() if "formula" in cp else ""
    if not name:
        raise Stop("[formula] summary_reference is not set in config/match_rate.ini. It names the most "
                   "complete standard built so far (D128), and it is what this tool scores against.")
    hit = [r for r in refs if r["name"] == name]
    if not hit:
        raise Stop(f"[formula] summary_reference {name!r} is not a [reference.*] in {CONFIG_INI.name}")
    ref = dict(hit[0])
    ref["version"] = cp["meta"].get("version", "").strip() if "meta" in cp else ""
    ref["decimals"] = int(cp["output"].get("decimals", "4")) if "output" in cp else 4
    ref["how"] = scored_sets[ref["scored_set"]]["how"]
    ref["symbol_of"] = {v: k for k, v in three.items()}          # letter -> three-letter, lowercased
    ref["three"], ref["names"] = three, names
    return ref


def three_letter(ref: dict, letter: str) -> str:
    return ref["symbol_of"].get(letter, letter).capitalize()


def alphabetical_letters(ref: dict) -> list[str]:
    """The twenty letters in alphabetical order of their IUPAC trivial name (the order labels and
    supplier sheets print: Ala, Arg, Asn, Asp, Cys, Glu, Gln, Gly, His ...). Read from the table,
    never listed here."""
    by_name = {v: k for k, v in ref["names"].items()}
    return sorted(AA, key=lambda a: by_name[a])


# --------------------------------------------------------------------------- the workbook in

def _cell(v) -> str:
    return "" if v is None else str(v).strip()


def read_sheet(path: Path) -> tuple[list[str], list[list]]:
    """(header, rows) of the sheet named `foods`, else the workbook's first sheet.

    Values come back as openpyxl read them — a number stays a number — so the output can write
    the author's own cells back unchanged.
    """
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        raise Stop(f"{_rel(path)} could not be opened as a workbook: {e}")
    ws = wb["foods"] if "foods" in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    rows = [r for r in rows if any(_cell(c) for c in r)]
    if not rows:
        return [], []
    header = [_cell(c) for c in rows[0]]
    while header and header[-1] == "":
        header.pop()
    if not header:
        raise Stop(f"{_rel(path)} sheet {ws.title!r}: the first row is empty; it must be the column headers")
    body = [list(r[:len(header)]) + [None] * max(0, len(header) - len(r)) for r in rows[1:]]
    return header, body


def amino_acid_columns(header: list[str], ref: dict, where: str) -> dict[int, str]:
    """{column index: letter} for the headers that are an amino acid; everything else passes through.

    A header is an amino acid only if it is exactly a three-letter symbol or a trivial name from the
    IUPAC-IUBMB table (case and surrounding space ignored). Nothing is named in code.
    """
    found: dict[int, str] = {}
    for i, name in enumerate(header):
        key = name.strip().lower()
        letter = ref["three"].get(key) or ref["names"].get(key)
        if letter:
            found[i] = letter
    seen: dict[str, str] = {}
    for i, letter in found.items():
        if letter in seen:
            raise Stop(f"{where}: columns {seen[letter]!r} and {header[i]!r} are both {three_letter(ref, letter)}")
        seen[letter] = header[i]
    return found


def _number(raw, where: str, field: str) -> float | None:
    if raw is None or str(raw).strip() == "":
        return None
    try:
        return float(str(raw).strip())
    except ValueError:
        raise Stop(f"{where}: {field} = {raw!r} is not a number")


# --------------------------------------------------------------------------- the imputation (M8)

def impute_unreported(values: dict[str, float | None], reference: dict[str, float],
                      scored: list[str]) -> tuple[dict[str, float], list[str], str]:
    """Fill a scored amino acid the row does not report at the reference's own share (ratio 1.0).

    Returns (the vector to score, the letters imputed, "" or the reason it cannot be scored).

    Arithmetically this is scoring over the amino acids the row DOES report, with both sides
    renormalised over that subset — the same number either way. Doing it as an imputation rather
    than a shrunken scored set is what keeps the walk in `steps` on the full scored set, so the
    sheet reads the same for a row of nine and a row of eight.

    An imputed amino acid can never become the limiting one on its own: the shares of both sides
    sum to one over the scored set, so at least one ratio is at or below 1.0 and it belongs to a
    reported amino acid unless every ratio is exactly 1.0.
    """
    reported = [a for a in scored if values.get(a) is not None]
    missing = [a for a in scored if values.get(a) is None]
    if not reported:
        return {}, missing, "the row reports none of the scored amino acids"
    r_total = sum(reference[a] for a in scored)
    missing_share = sum(reference[a] for a in missing) / r_total
    reported_sum = sum(values[a] for a in reported)
    if reported_sum <= 0:
        return {}, missing, "every scored amino acid the row reports is zero"
    scored_total = reported_sum / (1.0 - missing_share)
    out = {a: float(values[a]) for a in reported}
    for a in missing:
        out[a] = reference[a] / r_total * scored_total
    return out, missing, ""


# --------------------------------------------------------------------------- scoring one workbook

def score_workbook(path: Path, ref: dict) -> dict:
    """{header, rows, scored} for one input workbook; `scored` is one dict per row."""
    header, body = read_sheet(path)
    if not header:
        return {"header": [], "rows": [], "scored": []}
    where = _rel(path)
    absent = [c for c in REQUIRED_COLUMNS if c not in header]
    if absent:
        raise Stop(f"{where}: the sheet has no column(s) {', '.join(absent)}. Required: "
                   f"{', '.join(REQUIRED_COLUMNS)}; optional: protein_g; then the amino acids by "
                   f"three-letter symbol.")
    aa_cols = amino_acid_columns(header, ref, where)
    if not aa_cols:
        raise Stop(f"{where}: no column is an amino acid. Headers must be three-letter symbols "
                   f"(His, Ile, Leu ...) or trivial names (Histidine ...).")
    i_of = {c: header.index(c) for c in REQUIRED_COLUMNS}
    i_protein = header.index("protein_g") if "protein_g" in header else None
    scored_letters = ref["letters"]

    out = []
    for n, row in enumerate(body, 2):          # 2 = the first body row of the sheet, as Excel numbers it
        rw = f"{where} row {n}"
        label = _cell(row[i_of["label"]])
        if not label:
            raise Stop(f"{rw}: label is empty; it is the product's identity")
        basis = _cell(row[i_of["basis"]]).lower()
        if basis not in BASES:
            raise Stop(f"{rw} ({label}): basis = {_cell(row[i_of['basis']])!r}; it must be "
                       f"{' or '.join(BASES)} — what the amounts are per")
        grams = _number(row[i_of["grams_basis"]], rw, "grams_basis")
        if grams is None or grams <= 0:
            raise Stop(f"{rw} ({label}): grams_basis must be a number above zero — how many grams of "
                       f"{basis} the amounts belong to")
        protein = _number(row[i_protein], rw, "protein_g") if i_protein is not None else None
        values = {letter: _number(row[i], rw, header[i]) for i, letter in aa_cols.items()}
        values = {**{a: None for a in AA}, **values}

        vector, imputed, reason = impute_unreported(values, ref["values"], scored_letters)
        entry = {"label": label, "basis": basis, "grams_basis": grams, "protein": protein,
                 "imputed": imputed, "reason": reason, "result": None, "steps": None, "vector": vector}
        if not reason:
            res = match_rate(vector, ref["values"], scored_letters)
            if res["score"] is None:
                entry["reason"] = res["reason"]
            else:
                entry["result"] = res
                entry["steps"] = sheet_steps(vector, ref["values"], scored_letters, res)
        out.append(entry)
    return {"header": header, "rows": body, "scored": out}


# --------------------------------------------------------------------------- the workbook out

def _round(x, d: int):
    return None if x is None else round(float(x), d)


def calculation_row(e: dict, ref: dict) -> list:
    """The appended columns for one row, in CALC_COLUMNS order."""
    d = ref["decimals"]
    res, scored = e["result"], ref["letters"]
    imputed = ", ".join(three_letter(ref, a) for a in e["imputed"])
    set_text = ", ".join(three_letter(ref, a) for a in scored)
    if res is None:
        return [None, "", set_text, imputed, e["reason"], None, None, None, None]
    eaa = sum(e["vector"][a] for a in scored)
    protein, grams = e["protein"], e["grams_basis"]
    per_product = e["basis"] == "product"
    return [
        _round(100 * res["score"], d),
        "+".join(three_letter(ref, a) for a in res["limiting"]),
        set_text,
        imputed,
        "",
        _round(eaa, d),
        _round(100 * eaa / protein, d) if protein else None,
        _round(100 * eaa / grams, d) if per_product else None,
        _round(100 * protein / grams, d) if (per_product and protein is not None) else None,
    ]


def steps_table(scored: list[dict], ref: dict) -> tuple[list[str], list[list]]:
    """The spreadsheet's walk, one row per scored product (M5)."""
    d, letters = ref["decimals"], ref["letters"]
    sym = [three_letter(ref, a) for a in letters]
    header = (["label", "match_rate_percent"]
              + [f"{s}_amount" for s in sym]                                 # rows 12-20
              + ["eaa_total", "reference_eaa_total", "percent_of_reference_total"]   # 22, A22, 23
              + [f"step3_scaled_{s}" for s in sym]                           # rows 30-37
              + [f"step7_percent_of_reference_{s}" for s in sym]             # rows 100-107
              + ["step7_min", "limiting_amino_acid"]                         # row 108
              + [f"step10b_{s}" for s in sym]                                # rows 145-152
              + ["step10b_total_need_to_consume", "percent_wasted", "percent_utilized"])   # 153-156
    rows = []
    for e in scored:
        if e["result"] is None:
            continue
        s, res = e["steps"], e["result"]
        row = [e["label"], _round(100 * res["score"], d)]
        row += [_round(e["vector"][a], d) for a in letters]
        row += [_round(s["food_total"], d), _round(s["ref_total"], d), _round(s["pct_of_ref_total"], d)]
        row += [_round(s["scaled"][a], d) for a in letters]
        row += [_round(res["ratio"][a], d) for a in letters]
        row += [_round(res["score"], d), "+".join(three_letter(ref, a) for a in res["limiting"])]
        if s["final"] is None:                     # a zero score: Step 10b divides by it
            row += [None] * len(letters) + [None, None, _round(s["utilized"], d)]
        else:
            row += [_round(s["final"][a], d) for a in letters]
            row += [_round(s["need"], d), _round(s["wasted"], d), _round(s["utilized"], d)]
        rows.append(row)
    return header, rows


def reference_sheet(ref: dict, source: Path) -> list[list]:
    """The provenance and the reference's own values, as rows and columns. No comment lines."""
    d = ref["decimals"]
    total = sum(ref["values"][a] for a in ref["letters"])
    fields = [
        ("generated", common.iso_now()),
        ("generated_by", GENERATED_BY),
        ("match_rate_version", ref["version"] or "(unlabelled)"),
        ("input_workbook", source.name),
        ("input_workbook_sha256", sha256_path(source)),
        ("reference", ref["name"]),
        ("reference_label", ref["label"]),
        ("reference_read_from", ref["where"]),
        ("reference_sha256", sha256_path(ref["path"]) if ref["path"].exists() else ""),
        ("reference_decisions", ref["decisions"]),
        ("scored_set", ref["scored_set"]),
        ("scored_set_how", ref["how"]),
        ("scored_amino_acids", ", ".join(three_letter(ref, a) for a in ref["letters"])),
        ("config", f"{_rel(CONFIG_INI)} sha256 {sha256_path(CONFIG_INI)}"),
        ("amino_acid_symbols", f"{_rel(common.AMINO_ACID_SYMBOLS_INI)} sha256 "
                               f"{sha256_path(common.AMINO_ACID_SYMBOLS_INI)}"),
        ("formula", "Match Rate = the smallest of (food share / reference share) over the scored "
                    "amino acids, both shares over the scored set only (M1, D87)"),
        ("unreported_amino_acids", "imputed at the reference's own share, so their ratio is 1.0 "
                                   "(M8); the same number as scoring over the reported subset"),
    ]
    rows: list[list] = [["field", "value"]] + [[k, v] for k, v in fields]
    rows.append([])
    rows.append(["amino_acid", "symbol", "reference_amount", "reference_share_percent", "scored"])
    for a in alphabetical_letters(ref):
        value = ref["values"].get(a)          # a transcribed reference carries only its own rows
        share = 100 * value / total if (a in ref["letters"] and value is not None) else None
        rows.append([a, three_letter(ref, a), _round(value, 4),
                     _round(share, 4), "yes" if a in ref["letters"] else ""])
    return rows


def write_workbook(out_path: Path, book: dict, ref: dict, source: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "match_rate"
    ws.append(list(book["header"]) + CALC_COLUMNS)
    for raw, e in zip(book["rows"], book["scored"]):
        ws.append(list(raw) + calculation_row(e, ref))
    ws.freeze_panes = "A2"

    st = wb.create_sheet("steps")
    header, rows = steps_table(book["scored"], ref)
    st.append(header)
    for r in rows:
        st.append(r)
    st.freeze_panes = "A2"

    rs = wb.create_sheet("reference")
    for r in reference_sheet(ref, source):
        rs.append(r)

    for sheet in (ws, st, rs):
        for i in range(1, sheet.max_column + 1):
            width = max((len(str(sheet.cell(row=r, column=i).value or "")) for r in range(1, min(sheet.max_row, 40) + 1)),
                        default=10)
            sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def _is_ours(path: Path) -> bool:
    """True when `path` is a workbook this tool wrote — checked before overwriting it."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        if "reference" not in wb.sheetnames:
            return False
        for row in wb["reference"].iter_rows(values_only=True):
            if row and _cell(row[0]) == "generated_by":
                return _cell(row[1] if len(row) > 1 else "") == GENERATED_BY
        return False
    except Exception:
        return False
    finally:
        try:
            wb.close()
        except Exception:
            pass


# --------------------------------------------------------------------------- the blank sheet

def write_blank_sheet(path: Path, ref: dict) -> None:
    """A workbook with nothing but the column headers, written when scoring/ holds no input."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "foods"
    ws.append(REQUIRED_COLUMNS + ["protein_g", "source", "location", "note"]
              + [three_letter(ref, a) for a in alphabetical_letters(ref)])
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------- run

def input_workbooks(folder: Path) -> list[Path]:
    return sorted(p for p in folder.glob("*.xlsx")
                  if not p.name.startswith("~$") and not p.stem.endswith(SCORED_SUFFIX))


def run(log) -> list[Path]:
    folder = common.scoring_dir()
    folder.mkdir(parents=True, exist_ok=True)
    ref = load_summary_reference()
    log.info("reference %s (%s), scored on %s", ref["name"], ref["where"],
             ", ".join(three_letter(ref, a) for a in ref["letters"]))

    books = input_workbooks(folder)
    if not books:
        blank = folder / "blank_scoring_sheet.xlsx"
        if not blank.exists():
            write_blank_sheet(blank, ref)
            log.info("no workbook in %s; wrote %s with the column headers", _rel(folder), _rel(blank))
        else:
            log.info("no workbook to score in %s", _rel(folder))
        return []

    written = []
    for src in books:
        book = score_workbook(src, ref)
        if not book["scored"]:
            log.info("%s: no rows below the header; nothing scored", _rel(src))
            continue
        out = src.with_name(f"{src.stem}{SCORED_SUFFIX}.xlsx")
        if out.exists() and not _is_ours(out):
            raise Stop(f"{_rel(out)} exists and was not written by {GENERATED_BY}. Move or rename it; "
                       f"this tool will not overwrite a workbook it did not write.")
        write_workbook(out, book, ref, src)
        ok = [e for e in book["scored"] if e["result"] is not None]
        imputed = sum(1 for e in ok if e["imputed"])
        log.info("%s -> %s: %d row(s), %d scored, %d with an imputed amino acid",
                 _rel(src), _rel(out), len(book["scored"]), len(ok), imputed)
        for e in book["scored"]:
            if e["result"] is None:
                log.warning("  %s: not scored — %s", e["label"], e["reason"])
            else:
                log.info("  %s: %.2f %%, limiting %s%s", e["label"], 100 * e["result"]["score"],
                         "+".join(three_letter(ref, a) for a in e["result"]["limiting"]),
                         (" (imputed " + ", ".join(three_letter(ref, a) for a in e["imputed"]) + ")")
                         if e["imputed"] else "")
        written.append(out)
    return written


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.parse_args(argv)
    log = common.make_logger("score")
    run(log)
    return 0


if __name__ == "__main__":
    sys.exit(main())
