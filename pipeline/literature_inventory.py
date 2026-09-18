#!/usr/bin/env python3
"""
literature_inventory.py — offline stage: what is inside every literature file.

For every file recorded in data/literature/manifest.ini this stage re-hashes the
file on disk and compares it with the manifest (rule B0), then opens it without
changing it and records its structure:

  - a workbook (.xlsx): every sheet, its row and column counts, and the first
    rows of cells (header candidates);
  - an archive (.zip, .rar): every member with its size and SHA-256, and for
    every tabular member (.xlsx, .csv, .tsv, .txt) the same structure read
    from the bytes in memory — nothing is extracted to disk (handoff rule for
    the Deshmukh archives);
  - a document (.pdf, .docx): recorded by size and hash only; not opened.

Outputs
  outputs/literature_inventory/literature_files_and_members.tsv
      one row per file, archive member, and sheet: source_id, file_key, path,
      member, kind, bytes, sha256, sha256_matches_manifest, sheet, n_rows, n_cols
  outputs/literature_inventory/literature_header_rows.tsv
      one row per previewed sheet row: source_id, path, member, sheet,
      row_index (1-based, as a spreadsheet counts), cells (tab replaced by ' | ')
  outputs/literature_inventory/literature_inventory_summary.ini
  logs/literature_inventory_<UTC>.log

Why this exists
  The mass-fraction stage reads columns by the header strings named in
  config/mass_fraction_decisions.ini. Those strings are transcribed from the
  files' own headers, cited by sheet and row. This stage is the record of what
  those headers are, produced by code from the hashed files, so that the
  transcription can be checked against it and so that a reader without the
  files can see their shape.

Rules
  I1  Every manifest file must exist on disk and hash to the manifest's sha256;
      a mismatch or a missing file stops the run (B0).
  I2  Nothing is written under data/literature/. Archive members are read into
      memory; their hashes go into the output table and the log.
  I3  A .rar archive needs an extractor on this machine (WinRAR's UnRAR.exe,
      7-Zip's 7z.exe, unrar, unar, or bsdtar). The stage looks on PATH and in the
      default Windows install folders; if none is found it stops and says so.
      `--rar-tool <path>` overrides.
  I4  Cell values are recorded as text, truncated to --cell-chars characters;
      --preview-rows rows per sheet; at most --preview-cols cells per row.
      Row and column counts are exact.
  I5  Nothing is named, filtered, or interpreted. The stage records; it decides
      nothing.

Usage
  python pipeline/literature_inventory.py
  python pipeline/literature_inventory.py --only murgia_2021
  python pipeline/literature_inventory.py --rar-tool "C:\\Program Files\\WinRAR\\UnRAR.exe"
"""

from __future__ import annotations

import argparse
import hashlib
import io
import os
import shutil
import sys
import zipfile
from pathlib import Path

from pipeline import common

MANIFEST_INI = common.LITERATURE_MANIFEST_INI
OUT_DIR = common.LITERATURE_INVENTORY_DIR
FILES_TSV = OUT_DIR / "literature_files_and_members.tsv"
HEADERS_TSV = OUT_DIR / "literature_header_rows.tsv"
SUMMARY_INI = OUT_DIR / "literature_inventory_summary.ini"

TABULAR_EXT = {".xlsx", ".csv", ".tsv", ".txt"}
ARCHIVE_EXT = {".zip", ".rar"}
DOCUMENT_EXT = {".pdf", ".docx"}

FILES_HEADER = ["source_id", "file_key", "path", "member", "kind", "bytes", "sha256",
                "sha256_matches_manifest", "sheet", "n_rows", "n_cols"]
HEADERS_HEADER = ["source_id", "path", "member", "sheet", "row_index", "cells"]


# ----------------------------------------------------------------------------
# helpers
# ----------------------------------------------------------------------------

def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def cell_text(v, cell_chars: int) -> str:
    if v is None:
        return ""
    s = str(v).replace("\t", " ").replace("\n", " ").replace("\r", " ")
    return s if len(s) <= cell_chars else s[: cell_chars - 1] + "…"


def find_rar_tool(explicit: str | None, log) -> str | None:
    """Return a path to a RAR extractor, or None. Looks at --rar-tool, PATH, and
    the default Windows install folders. Tooling only; nothing biological."""
    candidates: list[str] = []
    if explicit:
        candidates.append(explicit)
    for name in ("unrar", "UnRAR", "7z", "7zz", "unar", "bsdtar"):
        found = shutil.which(name)
        if found:
            candidates.append(found)
    for env in ("ProgramFiles", "ProgramFiles(x86)", "ProgramW6432"):
        base = os.environ.get(env)
        if base:
            for rel in (r"WinRAR\UnRAR.exe", r"7-Zip\7z.exe"):
                p = Path(base) / rel
                if p.exists():
                    candidates.append(str(p))
    for c in candidates:
        if Path(c).exists():
            log.info("rar extractor: %s", c)
            return c
    return None


# ----------------------------------------------------------------------------
# inspectors — each returns (file_rows, header_rows)
# ----------------------------------------------------------------------------

def inspect_workbook(data: bytes, ctx: dict, opts) -> tuple[list[list], list[list]]:
    """Every sheet of an .xlsx from bytes in memory: exact row/column counts and
    the first rows as text. Read-only streaming; the workbook is never modified."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    file_rows, header_rows = [], []
    for ws in wb.worksheets:
        n_rows, n_cols = 0, 0
        for r_index, row in enumerate(ws.iter_rows(values_only=True), 1):
            n_rows = r_index
            # rightmost non-empty cell defines the row's width
            width = 0
            for j, v in enumerate(row, 1):
                if v is not None and str(v) != "":
                    width = j
            n_cols = max(n_cols, width)
            if r_index <= opts.preview_rows:
                cells = [cell_text(v, opts.cell_chars) for v in row[: opts.preview_cols]]
                header_rows.append([ctx["source_id"], ctx["path"], ctx["member"], ws.title,
                                    r_index, " | ".join(cells)])
        file_rows.append([ctx["source_id"], ctx["file_key"], ctx["path"], ctx["member"], "sheet",
                          ctx["bytes"], ctx["sha256"], ctx["match"], ws.title, n_rows, n_cols])
    wb.close()
    return file_rows, header_rows


def inspect_text_table(data: bytes, ctx: dict, opts) -> tuple[list[list], list[list]]:
    """A .csv/.tsv/.txt from bytes: line count, widest line by its delimiter, first lines."""
    text = data.decode("utf-8", errors="replace")
    lines = text.splitlines()
    delim = "\t" if ctx["member"].lower().endswith((".tsv", ".txt")) else ","
    n_cols = max((len(ln.split(delim)) for ln in lines if ln.strip()), default=0)
    header_rows = []
    for i, ln in enumerate(lines[: opts.preview_rows], 1):
        cells = [cell_text(c, opts.cell_chars) for c in ln.split(delim)[: opts.preview_cols]]
        header_rows.append([ctx["source_id"], ctx["path"], ctx["member"], "", i, " | ".join(cells)])
    file_rows = [[ctx["source_id"], ctx["file_key"], ctx["path"], ctx["member"], "text_table",
                  ctx["bytes"], ctx["sha256"], ctx["match"], "", len(lines), n_cols]]
    return file_rows, header_rows


def inspect_tabular_bytes(data: bytes, ctx: dict, opts) -> tuple[list[list], list[list]]:
    ext = Path(ctx["member"] or ctx["path"]).suffix.lower()
    if ext == ".xlsx":
        return inspect_workbook(data, ctx, opts)
    return inspect_text_table(data, ctx, opts)


def read_archive_members(path: Path, rar_tool: str | None) -> list[tuple[str, int, bytes]]:
    """Every file member of a .zip or .rar as (name, size, bytes), read in memory (I2, I3).
    Shared with mass_fractions.py, which reads archive members the same way."""
    ext = path.suffix.lower()
    members: list[tuple[str, int, bytes]] = []
    if ext == ".zip":
        with zipfile.ZipFile(path) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                members.append((info.filename, info.file_size, zf.read(info)))
    elif ext == ".rar":
        try:
            import rarfile
        except ImportError:
            raise SystemExit("[STOP] the `rarfile` package is not installed: pip install -r requirements.txt")
        if rar_tool is None:
            raise SystemExit("[STOP] no RAR extractor found (UnRAR.exe, 7z.exe, unrar, unar, or bsdtar). "
                             "Install one, add it to PATH, or pass --rar-tool <path>.")
        tool_name = Path(rar_tool).name.lower()
        if tool_name.startswith("7z"):
            rarfile.SEVENZIP_TOOL = rar_tool
        elif tool_name.startswith("unar"):
            rarfile.UNAR_TOOL = rar_tool
        elif tool_name.startswith("bsdtar"):
            rarfile.BSDTAR_TOOL = rar_tool
        else:
            rarfile.UNRAR_TOOL = rar_tool
        with rarfile.RarFile(str(path)) as rf:
            for info in rf.infolist():
                if info.isdir():
                    continue
                members.append((info.filename, info.file_size, rf.read(info)))
    else:
        raise SystemExit(f"[STOP] not an archive: {path}")
    return members


def inspect_archive(path: Path, ctx: dict, opts, rar_tool: str | None, log) -> tuple[list[list], list[list]]:
    """List every member with size and sha256; inspect tabular members from memory."""
    members = read_archive_members(path, rar_tool)
    file_rows, header_rows = [], []
    for name, size, data in members:
        m_sha = sha256_bytes(data)
        log.info("  member %s  %d bytes  sha256 %s", name, size, m_sha)
        mctx = dict(ctx, member=name, bytes=size, sha256=m_sha, match="")
        m_ext = Path(name).suffix.lower()
        if m_ext in TABULAR_EXT:
            fr, hr = inspect_tabular_bytes(data, mctx, opts)
            file_rows.append([ctx["source_id"], ctx["file_key"], ctx["path"], name, "archive_member",
                              size, m_sha, "", "", "", ""])
            file_rows += fr
            header_rows += hr
        else:
            file_rows.append([ctx["source_id"], ctx["file_key"], ctx["path"], name, "archive_member",
                              size, m_sha, "", "", "", ""])
    return file_rows, header_rows


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--only", help="source id, e.g. murgia_2021")
    ap.add_argument("--rar-tool", help="path to UnRAR.exe / 7z.exe / unrar / unar / bsdtar")
    ap.add_argument("--preview-rows", type=int, default=6)
    ap.add_argument("--preview-cols", type=int, default=40)
    ap.add_argument("--cell-chars", type=int, default=60)
    args = ap.parse_args(argv)

    log = common.make_logger("literature_inventory")
    if not MANIFEST_INI.exists():
        log.error("manifest not found: %s — run fetch_literature first", MANIFEST_INI)
        return 1
    manifest = common.read_ini(MANIFEST_INI)
    sections = [s for s in manifest.sections() if s.startswith("file.")]
    if args.only:
        sections = [s for s in sections if manifest[s].get("source_id") == args.only]
    log.info("manifest %s: %d file sections", MANIFEST_INI.relative_to(common.REPO_ROOT).as_posix(), len(sections))

    rar_tool = find_rar_tool(args.rar_tool, log)
    file_rows: list[list] = []
    header_rows: list[list] = []
    stops: list[str] = []
    n_archives = n_workbooks = n_documents = 0

    for sec in sections:
        m = manifest[sec]
        source_id = m.get("source_id", "")
        file_key = sec[len("file."):]
        rel = m.get("path", "")
        path = common.REPO_ROOT / rel
        expected = m.get("sha256", "").lower()
        if not path.exists():
            stops.append(f"{sec}: file missing on disk: {rel}")
            log.error(stops[-1])
            continue
        data = path.read_bytes()
        got = sha256_bytes(data)
        match = got == expected
        log.info("%s  %s  %d bytes  sha256 %s  manifest match: %s", sec, rel, len(data), got, match)
        if not match:
            stops.append(f"{sec}: sha256 on disk {got} != manifest {expected} (I1/B0)")
            log.error(stops[-1])
            continue
        ctx = {"source_id": source_id, "file_key": file_key, "path": rel, "member": "",
               "bytes": len(data), "sha256": got, "match": "true"}
        ext = path.suffix.lower()
        try:
            if ext in ARCHIVE_EXT:
                n_archives += 1
                file_rows.append([source_id, file_key, rel, "", "archive", len(data), got, "true", "", "", ""])
                fr, hr = inspect_archive(path, ctx, args, rar_tool, log)
                file_rows += fr
                header_rows += hr
            elif ext == ".xlsx":
                n_workbooks += 1
                file_rows.append([source_id, file_key, rel, "", "workbook", len(data), got, "true", "", "", ""])
                fr, hr = inspect_workbook(data, ctx, args)
                file_rows += fr
                header_rows += hr
            elif ext in TABULAR_EXT:
                fr, hr = inspect_text_table(data, ctx, args)
                file_rows += fr
                header_rows += hr
            elif ext in DOCUMENT_EXT:
                n_documents += 1
                file_rows.append([source_id, file_key, rel, "", "document_not_opened", len(data), got, "true", "", "", ""])
            else:
                file_rows.append([source_id, file_key, rel, "", "unknown_not_opened", len(data), got, "true", "", "", ""])
        except SystemExit as e:
            stops.append(f"{sec}: {e}")
            log.error(stops[-1])
        except Exception as e:  # a corrupt or unreadable file is a stop, not a silent skip
            stops.append(f"{sec}: could not open {rel}: {type(e).__name__}: {e}")
            log.error(stops[-1])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    header = [
        "GENERATED by literature_inventory.py. DO NOT EDIT BY HAND.",
        f"generated  = {common.iso_now()}",
        f"manifest   = {MANIFEST_INI.relative_to(common.REPO_ROOT).as_posix()}",
        f"preview    = {args.preview_rows} rows x {args.preview_cols} cells, {args.cell_chars} chars per cell (I4)",
    ]
    write_tsv_with_header(FILES_TSV, header, FILES_HEADER, file_rows)
    write_tsv_with_header(HEADERS_TSV, header, HEADERS_HEADER, header_rows)

    summ = common.new_ini()
    summ.add_section("result")
    summ["result"]["generated"] = common.iso_now()
    summ["result"]["manifest_files"] = str(len(sections))
    summ["result"]["archives"] = str(n_archives)
    summ["result"]["workbooks"] = str(n_workbooks)
    summ["result"]["documents_not_opened"] = str(n_documents)
    summ["result"]["sheets_and_members_listed"] = str(sum(1 for r in file_rows if r[4] in ("sheet", "archive_member", "text_table")))
    summ["result"]["rar_tool"] = rar_tool or "none found"
    summ["result"]["stops"] = str(len(stops))
    common.write_ini(summ, SUMMARY_INI, ["GENERATED by literature_inventory.py. DO NOT EDIT BY HAND."])

    log.info("files/members/sheets rows: %d; header rows: %d; stops: %d", len(file_rows), len(header_rows), len(stops))
    for s in stops:
        log.error("STOP: %s", s)
    return 1 if stops else 0


def write_tsv_with_header(path: Path, header_lines: list[str], columns: list[str], rows: list[list]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        for line in header_lines:
            fh.write(f"# {line}\n")
        fh.write("\t".join(columns) + "\n")
        for row in rows:
            fh.write("\t".join("" if v is None else str(v) for v in row) + "\n")


if __name__ == "__main__":
    sys.exit(main())
