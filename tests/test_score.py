"""The workbook scorer (`python run.py score`), on a synthetic repository.

The transcription test lives here too: the four foods of the author's spreadsheet
(Base_Match_Rate_Formula.xlsx, 2026-09-17, cells G156:J156) must come back out of a workbook with
the scores Excel cached, since this is the path a hand-entered food now takes (D129). The formula
itself is tested in tests/test_match.py; what is tested here is the workbook in, the workbook out,
and the imputation of an amino acid a label does not print (M8).

Nothing here reads a real standard, a real archive, or the author's own scoring/ folder.
"""
import pytest

from pipeline import common, match, score

SYMBOLS = {"A": ("Alanine", "Ala"), "C": ("Cysteine", "Cys"), "D": ("Aspartic acid", "Asp"),
           "E": ("Glutamic acid", "Glu"), "F": ("Phenylalanine", "Phe"), "G": ("Glycine", "Gly"),
           "H": ("Histidine", "His"), "I": ("Isoleucine", "Ile"), "K": ("Lysine", "Lys"),
           "L": ("Leucine", "Leu"), "M": ("Methionine", "Met"), "N": ("Asparagine", "Asn"),
           "P": ("Proline", "Pro"), "Q": ("Glutamine", "Gln"), "R": ("Arginine", "Arg"),
           "S": ("Serine", "Ser"), "T": ("Threonine", "Thr"), "V": ("Valine", "Val"),
           "W": ("Tryptophan", "Trp"), "Y": ("Tyrosine", "Tyr")}

# the spreadsheet: the reference it scored against (Gorissen 2018 Table 1) and its four foods
SHEET_REFERENCE = {"H": 2.8, "I": 3.4, "L": 6.3, "K": 6.6, "M": 1.7, "F": 3.8, "T": 2.9, "V": 4.3}
SHEET_FOODS = {
    "Rice": {"H": 2.83, "I": 3.41, "L": 6.57, "K": 2.79, "M": 3.20, "F": 4.18, "T": 3.16, "V": 3.74},
    "Chia": {"H": 0.531, "I": 0.801, "L": 1.371, "K": 0.97, "M": 0.588, "F": 1.016, "T": 0.709, "V": 0.95},
    "Pea": {"H": 1.99, "I": 4.05, "L": 6.80, "K": 6.11, "M": 0.85, "F": 3.86, "T": 3.05, "V": 4.28},
    "ChickBreastWhite": {"H": 0.717, "I": 1.219, "L": 1.732, "K": 1.962, "M": 0.639, "F": 0.916,
                         "T": 0.975, "V": 1.145},
}
SHEET_SCORES = {"Rice": 0.4498904709748083, "Chia": 0.673823005137884,
                "Pea": 0.5130687318489835, "ChickBreastWhite": 0.8238015780989283}
SHEET_LIMITING = {"Rice": "Lys", "Chia": "Lys", "Pea": "Met", "ChickBreastWhite": "Phe"}
EIGHT = ["T", "M", "F", "H", "K", "V", "I", "L"]
NINE_STANDARD = {"A": 6.5, "C": 1.2, "D": 5.5, "E": 12.0, "F": 4.8, "G": 4.1, "H": 3.1, "I": 6.5,
                 "K": 11.3, "L": 10.0, "M": 3.2, "N": 3.5, "P": 4.5, "Q": 4.5, "R": 8.2, "S": 5.7,
                 "T": 6.1, "V": 7.1, "W": 1.2, "Y": 4.6}


class _Log:
    def info(self, *a):
        pass

    def warning(self, *a):
        pass

    def error(self, *a):
        pass


# --------------------------------------------------------------------------- fixtures

def _write_config(path, summary="eight_reference"):
    path.write_text(
        "[meta]\nversion = test\n\n[formula]\nprimary_reference = nine_reference\n"
        f"summary_reference = {summary}\n\n"
        "[scored_set.nine]\ndecision = D86\nkind = fao_headings_reduced\nfrom_file = config/fao.ini\n"
        "group_scored_as.SAA = Met\ngroup_scored_as.AAA = Phe\n\n"
        "[scored_set.eight]\nkind = rows_named_in_file\nfrom_file = config/gorissen.ini\n"
        "from_section = values.human_muscle.location\nfrom_key = essential_rows\n\n"
        "[reference.nine_reference]\nfile = outputs/standard/nine.tsv\ncolumn = standard\n"
        "scored_set = nine\nlabel = synthetic nine-amino-acid reference\ndecisions = test\n\n"
        "[reference.eight_reference]\nconfig = config/gorissen.ini\nsection = values.human_muscle\n"
        "scored_set = eight\nlabel = synthetic eight-amino-acid reference\ndecisions = test\n\n"
        "[foods.usda]\nfile = outputs/usda/foods.tsv\n\n[output]\ndecimals = 4\n", encoding="utf-8")


@pytest.fixture
def repo(tmp_path, monkeypatch):
    """A repository with the two references the spreadsheet work uses and an empty scoring/ folder."""
    (tmp_path / "config").mkdir()
    (tmp_path / "outputs" / "standard").mkdir(parents=True)
    (tmp_path / "data" / "iupac").mkdir(parents=True)

    cp = common.new_ini()
    for a in common.AMINO_ACIDS:
        cp.add_section(a)
        cp[a]["trivial_name"] = SYMBOLS[a][0]
        cp[a]["three_letter"] = SYMBOLS[a][1]
    common.write_ini(cp, tmp_path / "data" / "iupac" / "symbols.ini", ["synthetic"])

    (tmp_path / "config" / "fao.ini").write_text(
        "[indispensable_amino_acids]\nheadings = His, Ile, Leu, Lys, SAA, AAA, Thr, Trp, Val\n"
        "location = synthetic\n\n[group.SAA]\nmembers = Met, Cys\n\n[group.AAA]\nmembers = Phe, Tyr\n",
        encoding="utf-8")

    values = {**{SYMBOLS[a][0]: SHEET_REFERENCE[a] for a in SHEET_REFERENCE},
              "Serine": 2.3, "Glycine": 3.1, "Glutamic acid": 13.1, "Alanine": 4.1,
              "Tyrosine": 2.0, "Arginine": 4.4}
    body = "[values.human_muscle]\n" + "".join(f"{k} = {v}\n" for k, v in values.items())
    body += "\n[values.human_muscle.location]\nessential_rows = " + \
            ", ".join(SYMBOLS[a][0] for a in EIGHT) + "\n"
    (tmp_path / "config" / "gorissen.ini").write_text(body, encoding="utf-8")

    lines = ["# synthetic", "\t".join(["amino_acid", "three_letter", "name", "standard"])]
    for a in common.AMINO_ACIDS:
        lines.append("\t".join([a, SYMBOLS[a][1], SYMBOLS[a][0], f"{NINE_STANDARD[a]:.4f}"]))
    (tmp_path / "outputs" / "standard" / "nine.tsv").write_text("\n".join(lines) + "\n", encoding="utf-8")

    _write_config(tmp_path / "config" / "match_rate.ini")
    monkeypatch.setattr(common, "REPO_ROOT", tmp_path)
    monkeypatch.setattr(common, "OUTPUTS_DIR", tmp_path / "outputs")
    monkeypatch.setattr(common, "AMINO_ACID_SYMBOLS_INI", tmp_path / "data" / "iupac" / "symbols.ini")
    monkeypatch.setattr(score, "CONFIG_INI", tmp_path / "config" / "match_rate.ini")
    monkeypatch.setattr(match, "CONFIG_INI", tmp_path / "config" / "match_rate.ini")
    assert common.scoring_dir().is_relative_to(tmp_path), "the fixture did not take effect"
    (tmp_path / "scoring").mkdir()
    return tmp_path


def write_workbook(path, header, rows, sheet="foods"):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def read_sheet(path, name):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    names = wb.sheetnames
    wb.close()
    return names, rows


def by_label(path, name="match_rate"):
    _, rows = read_sheet(path, name)
    header = [c for c in rows[0] if c is not None]
    return header, {r[0]: dict(zip(header, r)) for r in rows[1:] if r and r[0] is not None}


def sheet_foods_workbook(path, ref_letters="HILKMFTV", extra_columns=()):
    header = ["label", "source", "basis", "grams_basis", "protein_g"] + \
             [SYMBOLS[a][1] for a in ref_letters] + list(extra_columns)
    rows = [[n, "the author's spreadsheet, 2026-09-17", "product", 100, None]
            + [SHEET_FOODS[n][a] for a in ref_letters] + [""] * len(extra_columns)
            for n in SHEET_FOODS]
    write_workbook(path, header, rows)
    return header


# --------------------------------------------------------------------------- the transcription (D87, D129)

def test_the_spreadsheets_four_foods_score_through_a_workbook(repo):
    """The four scores Excel cached, out of the workbook path rather than the deleted CSV."""
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    out = score.run(_Log())
    assert [p.name for p in out] == ["my_foods_scored.xlsx"]
    _, rows = by_label(out[0])
    for name in SHEET_FOODS:
        assert rows[name]["match_rate_percent"] == pytest.approx(100 * SHEET_SCORES[name], abs=5e-4), name
        assert rows[name]["limiting_amino_acid"] == SHEET_LIMITING[name], name


def test_the_steps_sheet_reproduces_the_spreadsheets_cells(repo):
    """Pea, column I: I22 = 30.99, I23 = 0.97453, I153 = 61.98, I155 = 0.48693, I156 = 0.51307."""
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    out = score.run(_Log())
    _, rows = by_label(out[0], "steps")
    pea = rows["Pea"]
    assert pea["eaa_total"] == pytest.approx(30.99, abs=1e-3)
    assert pea["reference_eaa_total"] == pytest.approx(31.8, abs=1e-3)
    assert pea["percent_of_reference_total"] == pytest.approx(30.99 / 31.8, abs=1e-4)
    assert pea["step7_min"] == pytest.approx(SHEET_SCORES["Pea"], abs=1e-4)
    assert pea["limiting_amino_acid"] == "Met"
    assert pea["step10b_Met"] == pytest.approx(1.7, abs=1e-3)       # the limiting one meets the reference
    assert pea["step10b_total_need_to_consume"] == pytest.approx(31.8 / SHEET_SCORES["Pea"], abs=1e-2)
    assert pea["percent_wasted"] + pea["percent_utilized"] == pytest.approx(1.0, abs=1e-3)


# --------------------------------------------------------------------------- M9: the author's rows, untouched

def test_the_input_rows_come_back_verbatim_with_the_calculations_appended(repo):
    header = sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx", extra_columns=("note",))
    out = score.run(_Log())
    cols, rows = by_label(out[0])
    assert cols[:len(header)] == header, "the author's columns moved or changed"
    assert cols[len(header):] == score.CALC_COLUMNS
    for name, food in SHEET_FOODS.items():
        assert rows[name]["source"] == "the author's spreadsheet, 2026-09-17"
        for letter, value in food.items():
            assert rows[name][SYMBOLS[letter][1]] == pytest.approx(value), (name, letter)


def test_the_input_workbook_is_not_written_to(repo):
    src = repo / "scoring" / "my_foods.xlsx"
    sheet_foods_workbook(src)
    before = src.read_bytes()
    score.run(_Log())
    assert src.read_bytes() == before


def test_three_sheets_and_no_comment_rows(repo):
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    out = score.run(_Log())
    names, _ = read_sheet(out[0], "match_rate")
    assert names == ["match_rate", "steps", "reference"]
    for sheet in names:
        _, rows = read_sheet(out[0], sheet)
        for r in rows:
            first = "" if not r or r[0] is None else str(r[0])
            assert not first.startswith("#"), f"{sheet}: a comment row"


def test_the_reference_sheet_carries_the_provenance_and_the_values(repo):
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    out = score.run(_Log())
    _, rows = read_sheet(out[0], "reference")
    fields = {r[0]: (r[1] if len(r) > 1 else None) for r in rows if r and r[0]}
    assert fields["generated_by"] == score.GENERATED_BY
    assert fields["reference"] == "eight_reference"
    assert len(str(fields["input_workbook_sha256"])) == 64
    assert "sha256" in str(fields["config"])
    assert fields["scored_amino_acids"] == ", ".join(SYMBOLS[a][1] for a in EIGHT)
    table = [r for r in rows if r and r[0] in common.AMINO_ACIDS]
    assert len(table) == 20, "every amino acid is listed, scored or not"
    scored = [r[1] for r in table if r[4] == "yes"]
    assert set(scored) == {SYMBOLS[a][1] for a in EIGHT}


def test_the_amino_acids_are_alphabetical_by_trivial_name(repo):
    """The order labels and supplier sheets print — so Asn before Asp, and Glu before Gln."""
    ref = score.load_summary_reference()
    order = [score.three_letter(ref, a) for a in score.alphabetical_letters(ref)]
    assert order[:9] == ["Ala", "Arg", "Asn", "Asp", "Cys", "Glu", "Gln", "Gly", "His"]
    assert order[-3:] == ["Trp", "Tyr", "Val"]


# --------------------------------------------------------------------------- M8: the imputation

def test_imputing_an_unreported_amino_acid_is_scoring_the_reported_subset(repo):
    """The identity the rule rests on: filling the missing one at the reference's own share gives
    the same number as renormalising both sides over the amino acids the row does report."""
    _write_config(repo / "config" / "match_rate.ini", summary="nine_reference")
    ref = score.load_summary_reference()
    for name, food in SHEET_FOODS.items():           # the sheet's foods carry no tryptophan
        values = {**{a: None for a in common.AMINO_ACIDS}, **food}
        vector, imputed, reason = score.impute_unreported(values, ref["values"], ref["letters"])
        assert imputed == ["W"] and not reason, name
        got = match.match_rate(vector, ref["values"], ref["letters"])
        reported = [a for a in ref["letters"] if a != "W"]
        direct = match.match_rate({a: food[a] for a in reported},
                                  {a: ref["values"][a] for a in reported}, reported)
        assert got["score"] == pytest.approx(direct["score"], abs=1e-12), name
        assert got["ratio"]["W"] == pytest.approx(1.0, abs=1e-12), name
        assert "W" not in got["limiting"], f"{name}: an imputed amino acid became the limiting one"


def test_the_imputed_amino_acids_are_named_on_the_row(repo):
    _write_config(repo / "config" / "match_rate.ini", summary="nine_reference")
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    out = score.run(_Log())
    _, rows = by_label(out[0])
    for name in SHEET_FOODS:
        assert rows[name]["imputed_at_reference_share"] == "Trp", name
        assert rows[name]["scored_amino_acids"].count(",") == 8, "nine scored"


def test_a_blank_is_imputed_and_a_zero_is_scored(repo):
    """M3 and M8 are different rules: blank means the source did not report it; 0 means it measured zero."""
    _write_config(repo / "config" / "match_rate.ini", summary="nine_reference")
    ref = score.load_summary_reference()
    header = ["label", "basis", "grams_basis"] + [SYMBOLS[a][1] for a in ref["letters"]]
    write_workbook(repo / "scoring" / "two.xlsx", header, [
        ["blank lysine", "product", 100] + [None if a == "K" else 1.0 for a in ref["letters"]],
        ["zero lysine", "product", 100] + [0.0 if a == "K" else 1.0 for a in ref["letters"]],
    ])
    out = score.run(_Log())
    _, rows = by_label(out[0])
    assert rows["zero lysine"]["match_rate_percent"] == 0
    assert rows["zero lysine"]["limiting_amino_acid"] == "Lys"
    assert rows["blank lysine"]["match_rate_percent"] > 0
    assert rows["blank lysine"]["imputed_at_reference_share"] == "Lys"


def test_a_row_that_reports_nothing_is_not_scored_and_says_why(repo):
    header = ["label", "basis", "grams_basis", "Lys", "Leu"]
    write_workbook(repo / "scoring" / "thin.xlsx", header, [
        ["two of eight", "product", 100, 2.0, 3.0],
        ["reports none", "product", 100, None, None],
        ["all zero", "product", 100, 0, 0],
    ])
    out = score.run(_Log())
    _, rows = by_label(out[0])
    assert rows["two of eight"]["match_rate_percent"] is not None
    assert rows["reports none"]["match_rate_percent"] is None
    assert "none of the scored" in rows["reports none"]["not_scored_reason"]
    assert rows["all zero"]["match_rate_percent"] is None
    assert rows["all zero"]["not_scored_reason"]
    _, steps = by_label(out[0], "steps")
    assert list(steps) == ["two of eight"], "the steps sheet carries only what was scored"


# --------------------------------------------------------------------------- the basis

def test_the_basis_changes_the_density_columns_and_never_the_score(repo):
    ref = score.load_summary_reference()
    header = ["label", "basis", "grams_basis", "protein_g"] + [SYMBOLS[a][1] for a in ref["letters"]]
    n = len(ref["letters"])
    write_workbook(repo / "scoring" / "basis.xlsx", header, [
        ["per product", "product", 250, 60] + [2.0] * n,
        ["per protein", "protein", 100, 100] + [2.0] * n,
    ])
    out = score.run(_Log())
    _, rows = by_label(out[0])
    assert rows["per product"]["match_rate_percent"] == pytest.approx(rows["per protein"]["match_rate_percent"])
    assert rows["per product"]["eaa_g_per_100g_product"] == pytest.approx(100 * 2.0 * n / 250)
    assert rows["per product"]["protein_percent_of_product"] == pytest.approx(100 * 60 / 250)
    assert rows["per protein"]["eaa_g_per_100g_product"] is None
    assert rows["per protein"]["protein_percent_of_product"] is None
    assert rows["per protein"]["eaa_percent_of_protein"] == pytest.approx(2.0 * n)


def test_protein_is_optional(repo):
    header = ["label", "basis", "grams_basis", "Lys", "Leu"]
    write_workbook(repo / "scoring" / "noprotein.xlsx", header, [["x", "product", 100, 2.0, 3.0]])
    out = score.run(_Log())
    _, rows = by_label(out[0])
    assert rows["x"]["match_rate_percent"] is not None
    assert rows["x"]["eaa_percent_of_protein"] is None


# --------------------------------------------------------------------------- what stops

@pytest.mark.parametrize("row, message", [
    (["x", "per serving", 100, 1.0], "basis"),
    (["x", "product", 0, 1.0], "grams_basis"),
    (["x", "product", "many", 1.0], "grams_basis"),
    (["", "product", 100, 1.0], "label"),
])
def test_a_row_the_tool_cannot_read_stops_it(repo, row, message):
    write_workbook(repo / "scoring" / "bad.xlsx", ["label", "basis", "grams_basis", "Lys"], [row])
    with pytest.raises(SystemExit, match=message):
        score.run(_Log())


def test_a_sheet_without_the_required_columns_stops(repo):
    write_workbook(repo / "scoring" / "bad.xlsx", ["label", "Lys"], [["x", 1.0]])
    with pytest.raises(SystemExit, match="grams_basis"):
        score.run(_Log())


def test_a_sheet_with_no_amino_acid_column_stops(repo):
    write_workbook(repo / "scoring" / "bad.xlsx", ["label", "basis", "grams_basis", "protein_g"],
                   [["x", "product", 100, 20]])
    with pytest.raises(SystemExit, match="no column is an amino acid"):
        score.run(_Log())


def test_the_same_amino_acid_in_two_columns_stops(repo):
    write_workbook(repo / "scoring" / "bad.xlsx",
                   ["label", "basis", "grams_basis", "Lys", "Lysine"], [["x", "product", 100, 1.0, 1.0]])
    with pytest.raises(SystemExit, match="Lys"):
        score.run(_Log())


def test_a_summary_reference_that_is_not_a_reference_stops(repo):
    _write_config(repo / "config" / "match_rate.ini", summary="nonexistent")
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    with pytest.raises(SystemExit, match="summary_reference"):
        score.run(_Log())


def test_a_file_that_is_not_a_workbook_stops(repo):
    (repo / "scoring" / "notreally.xlsx").write_text("not a workbook", encoding="utf-8")
    with pytest.raises(SystemExit, match="could not be opened"):
        score.run(_Log())


# --------------------------------------------------------------------------- the folder (M7)

def test_a_scored_workbook_is_never_taken_as_an_input(repo):
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    score.run(_Log())
    assert (repo / "scoring" / "my_foods_scored.xlsx").exists()
    assert [p.name for p in score.input_workbooks(repo / "scoring")] == ["my_foods.xlsx"]
    again = score.run(_Log())
    assert [p.name for p in again] == ["my_foods_scored.xlsx"], "a rerun rescores the input, not its output"


def test_a_workbook_the_tool_did_not_write_is_not_overwritten(repo):
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    write_workbook(repo / "scoring" / "my_foods_scored.xlsx", ["mine"], [["not the tool's"]])
    with pytest.raises(SystemExit, match="will not overwrite"):
        score.run(_Log())


def test_an_empty_folder_gets_a_sheet_with_the_column_headers(repo):
    out = score.run(_Log())
    assert out == []
    blank = repo / "scoring" / "blank_scoring_sheet.xlsx"
    assert blank.exists()
    _, rows = read_sheet(blank, "foods")
    header = [c for c in rows[0] if c is not None]
    assert header[:3] == score.REQUIRED_COLUMNS
    assert header[-20:] == [SYMBOLS[a][1] for a in
                            sorted(common.AMINO_ACIDS, key=lambda x: SYMBOLS[x][0].lower())]
    assert score.run(_Log()) == [], "a sheet with no rows below the header produces nothing"


def test_the_scoring_folder_follows_a_redirected_repo_root(tmp_path, monkeypatch):
    """D91: a path derived from a root that can move is resolved when it is used."""
    monkeypatch.setattr(common, "REPO_ROOT", tmp_path / "elsewhere")
    assert common.scoring_dir() == tmp_path / "elsewhere" / "scoring"


def test_the_scorer_is_not_a_pipeline_command(repo):
    """M7: tooling. Nothing depends on it, so it has no place in the run order and writes
    nothing under outputs/."""
    assert "score" not in common.COMMAND_ORDER
    sheet_foods_workbook(repo / "scoring" / "my_foods.xlsx")
    score.run(_Log())
    written = [p for p in (repo / "outputs").rglob("*") if p.is_file()]
    assert all("scor" not in p.name for p in written), [p.name for p in written]
